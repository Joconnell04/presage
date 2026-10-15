"""Scratch paper: run the client's own model instead of imitating it.

A client asks "what if we charged 1% of GMV". No cell holds that number, and the copilot
must never do the arithmetic in its head. So this module compiles the workbook with pycel,
sets the input cell, re-evaluates, and hands back before/after with citations.

Six pieces:

  ModelGraph      compile once per workbook, evaluate, what-if with a guaranteed restore,
                  and a record of every cell pycel cannot evaluate (`partial`)
  discover        the lever table: the inputs a formula reads, numeric AND text, with the
                  labelled summary-sheet outputs each one moves
  sensitivity     a few round values per lever plus pairwise combinations of the top ones,
                  precomputed, so the instant path can answer
  solve / trace   goal seek by bisection, and the formula chain from a lever to a moved cell
  eval_expr       the `calc` grammar: a hand-written AST walker, never `eval`
  parse/format    "1%", "$4.2M", "1.2x", "12,500" in; "$9.59M", "0.85%" out

Scenario selectors ARE levers now. A text cell like `Summary Scenarios!B5 = "Premium"`
switches the whole model between named cases; its "grid" is the values the workbook allows
(the cell's data validation list, else the literal strings formulas compare it against), and
`whatif` takes one of them in place of a number.

== Index methods (presage/ingest/__init__.py). The platform owner adds TOOLS entries for the
   starred ones in presage/brain.py; the rest already have one. ==

  levers(file=None) -> [{file, sheet, cell, label, header, value, fmt, choices, aliases,
                         outputs: [{cell, label, fmt, value}], active, reads, reach}]
      `fmt` is pct | money | count | number | text. A text lever carries `choices`, the
      values the workbook allows; every other lever has `choices: []`.
  levers_block(limit=40) -> str            the LEVERS block for the system prompt
  whatif(file, sets, watch=None) -> {file, changes, outputs, partial, ms}
      `sets` values may be numbers, spoken numbers ("1%", "$4.2M"), or, for a text lever,
      one of its `choices` (matched loosely: "downside", "flat dues" both resolve).
      Several cells in one call is the multi-lever case and always was.
  calc(expr, inputs) -> {value, expr, inputs}

  * goalseek(file, target, value, lever, lo=None, hi=None) -> {lever, label, value, value_s,
        was, was_s, target, target_label, want, achieved, achieved_s, monotone, tries, ms}
        "what fee gets us to $10M": bisection on the compiled graph. Auto-brackets around
        the lever's current value; `error` when the target cannot be reached in range.
  * explain(file, cell, sets=None, hops=6) -> {cell, label, was, now, chain: [{cell, label,
        formula, was, now, delta}], lever, moved, ms}
        "why did costs move": the formula chain from the changed lever down to the cell,
        3 to 6 hops, with each intermediate value before and after.
  * compare(file, cases, watch=None, labels=None) -> {cases: [{label, sets, outputs, partial}],
        rows: [{cell, label, fmt, values, best, worst, spread}], ms}
        "premium vs market at 1%": several whole cases in one call, aligned by output cell.
  * sweep(file, cell, values, watch=None) -> {cell, label, points: [{value, value_s, outputs}],
        rows: [{cell, label, fmt, values}], ms}
        "give me year five revenue at 1, 1.25 and 1.5 percent": one lever, many values.

Threading: pycel's compiler is mutable shared state, so every workbook has one lock and
every public method that touches the graph holds it. All of it is sync and CPU-bound; the
server calls it through `asyncio.to_thread` so the event loop never waits.
"""

from __future__ import annotations

import ast
import logging
import math
import re
import threading
import time
from pathlib import Path

PLUGIN = "presage.ingest.xlfuncs"

# A sheet whose name looks like the place a partner would read a number off.
SUMMARY_SHEET_RE = re.compile(
    r"summary|overview|dashboard|recommendation|p\s*&\s*l|profit\s+and\s+loss|"
    r"valuation|output", re.I)

MAX_LEVERS = 250            # per workbook, ranked; the prompt block takes fewer still
MAX_OUTPUTS = 8             # per lever, nearest first
MAX_PER_LABEL = 2           # B33/C33/D33 "Year 5 penetration" are one lever three times
MAX_GRID_PER_LEVER = 10
MAX_GRID_RECALCS = 3000     # per workbook (about 3 ms a point on a 3,500-formula model)
MAX_RANGE_CELLS = 4000      # a formula reading a whole column is not evidence of a lever
BFS_NODE_BUDGET = 60000     # cap the downstream walk per lever
BFS_KEEP_CAP = 160          # raw output hits to collect before picking the best 8
INPUT_SHEET_RE = re.compile(r"assumption|input|driver|lever|control", re.I)
SCHEMA = 3                  # lever-table shape; a cached table from an older one is rebuilt
MAX_CHOICES = 12            # allowed values of a text lever
SECTION_UP = 15             # rows to look up for the block title a lever sits under
MAX_ALIASES = 8             # spoken names per lever: off the sheet, off the section, and the
                            # deep pass's. Three at four dropped the one name that told two
                            # 0.85% "platform fee" cells apart (E36).
COMBO_LEVERS = 10           # top levers paired up in the grid
MAX_COMBO_POINTS = 240      # pairwise recalcs per workbook (about 1 s)
GOALSEEK_TRIES = 60
EXPLAIN_HOPS = 6

log = logging.getLogger(__name__)


# --------------------------------------------------------------------------- #
# cell references
# --------------------------------------------------------------------------- #
_SHEET_CELL = re.compile(
    r"^\s*(?:'(?P<q>[^']+)'|(?P<s>[^'!]+))!(?P<c>\$?[A-Za-z]{1,3}\$?[0-9]{1,7})\s*$")
_CELL_ONLY = re.compile(r"^\s*(?P<c>\$?[A-Za-z]{1,3}\$?[0-9]{1,7})\s*$")

# References inside a formula. The lookbehind stops `1E5` and `LOG10` from reading as cells.
_REF_IN_FORMULA = re.compile(
    r"(?<![A-Za-z0-9_.$!])"
    r"(?:(?:'(?P<q>[^']{1,120})'|(?P<s>[A-Za-z_][A-Za-z0-9_.]{0,60}))!)?"
    r"\$?(?P<c1>[A-Z]{1,3})\$?(?P<r1>[1-9][0-9]{0,6})"
    r"(?::\$?(?P<c2>[A-Z]{1,3})\$?(?P<r2>[1-9][0-9]{0,6}))?"
    r"(?![A-Za-z0-9_(])")
_STRINGS = re.compile(r'"[^"]*"')


def col_num(letters: str) -> int:
    n = 0
    for ch in letters.upper():
        n = n * 26 + (ord(ch) - 64)
    return n


def col_letters(n: int) -> str:
    s = ""
    while n > 0:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def split_ref(ref: str, default_sheet: str = "") -> tuple[str, str]:
    """'Subscriber Fee Model'!B48 -> ('Subscriber Fee Model', 'B48'). Raises on nonsense."""
    text = str(ref or "").strip()
    m = _SHEET_CELL.match(text)
    if m:
        return (m.group("q") or m.group("s") or "").strip(), m.group("c").replace("$", "").upper()
    m = _CELL_ONLY.match(text)
    if m:
        if not default_sheet:
            raise ValueError(f"cell reference {ref!r} has no sheet (use 'Sheet!A1')")
        return default_sheet, m.group("c").replace("$", "").upper()
    raise ValueError(f"{ref!r} is not a cell reference like Sheet!A1")


def canon(sheet: str, cell: str) -> str:
    """pycel's own key format: unquoted sheet, no dollars, upper case."""
    return f"{sheet}!{cell.replace('$', '').upper()}"


def pretty(sheet: str, cell: str) -> str:
    """How a citation reads on the page: 'Subscriber Fee Model › B48'."""
    return f"{sheet} \u203a {cell}"


def refs_in_formula(formula: str, sheet: str) -> set[tuple[str, str]]:
    """Every cell a formula reads, ranges expanded, as {(sheet, 'B48')}."""
    body = _STRINGS.sub(" ", str(formula or ""))
    out: set[tuple[str, str]] = set()
    for m in _REF_IN_FORMULA.finditer(body):
        sh = (m.group("q") or m.group("s") or sheet).strip()
        if not sh or "[" in sh:                    # linked workbook: not ours to move
            continue
        c1, r1 = m.group("c1"), int(m.group("r1"))
        if m.group("c2") is None:
            out.add((sh, f"{c1}{r1}"))
            continue
        c2, r2 = m.group("c2"), int(m.group("r2"))
        lo_c, hi_c = sorted((col_num(c1), col_num(c2)))
        lo_r, hi_r = sorted((r1, r2))
        if (hi_c - lo_c + 1) * (hi_r - lo_r + 1) > MAX_RANGE_CELLS:
            continue
        for c in range(lo_c, hi_c + 1):
            letter = col_letters(c)
            for r in range(lo_r, hi_r + 1):
                out.add((sh, f"{letter}{r}"))
    return out


# --------------------------------------------------------------------------- #
# values in and out
# --------------------------------------------------------------------------- #
_VALUE = re.compile(
    r"^(?P<open>\()?\s*(?P<sign>[-+])?\s*(?P<cur>[$£€])?\s*"
    r"(?P<num>\d[\d,]*(?:\.\d+)?|\.\d+)\s*"
    r"(?P<suf>%|bps|bp|mm|bn|k|m|b|x)?\s*(?P<close>\))?$", re.I)
_SUFFIX = {"k": 1e3, "m": 1e6, "mm": 1e6, "b": 1e9, "bn": 1e9}


def parse_value(value) -> float:
    """'1%' -> 0.01, '$4.2M' -> 4200000.0, '1.2x' -> 1.2, '(1,200)' -> -1200.0."""
    if isinstance(value, bool):
        raise ValueError(f"{value!r} is not a number")
    if isinstance(value, (int, float)):
        return float(value)
    m = _VALUE.match(str(value or "").strip())
    if not m:
        raise ValueError(f"{value!r} is not a number I can read "
                         "(try 1%, 0.01, $4.2M, 1.2x or 12,500)")
    num = float(m.group("num").replace(",", ""))
    suf = (m.group("suf") or "").lower()
    if suf == "%":
        num /= 100.0
    elif suf in ("bps", "bp"):
        num /= 10000.0
    elif suf in _SUFFIX:
        num *= _SUFFIX[suf]
    if m.group("sign") == "-" or (m.group("open") and m.group("close")):
        num = -num
    return num


def fmt_kind(number_format: str, value=None) -> str:
    """A cell's number format -> pct | money | count | number."""
    nf = str(number_format or "").split(";")[0]
    if "%" in nf:
        return "pct"
    if re.search(r"\[\$|\\?[$£€¥]", nf):
        return "money"
    digits = re.sub(r"\\.|\".*?\"|\[[^\]]*\]", "", nf)
    if re.search(r"[#0]", digits):
        return "number" if "." in digits else "count"
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return "count" if float(value) == int(value) and abs(value) >= 1 else "number"
    return "number"


def _trim(x: float, places: int = 4) -> str:
    s = f"{x:,.{places}f}".rstrip("0").rstrip(".")
    return s or "0"


def fmt_value(value, kind: str = "number") -> str:
    """The way a number is spoken back: '0.85%', '$9.59M', '12,500', '1.2'."""
    if value is None:
        return "n/a"
    if isinstance(value, str):
        return value
    v = float(value)
    if kind == "pct":
        pct = v * 100.0
        return f"{_trim(pct, 2 if abs(pct) < 10 else 1)}%"
    if kind == "money":
        sign = "-" if v < 0 else ""
        a = abs(v)
        if a >= 1e6:
            return f"{sign}${_trim(a / 1e6, 2)}M"
        if a >= 1e4:
            return f"{sign}${a:,.0f}"
        if a >= 100 or a == int(a):
            return f"{sign}${a:,.0f}"
        return f"{sign}${a:,.2f}"
    if kind == "count":
        return f"{v:,.0f}" if abs(v - round(v)) < 1e-9 else _trim(v, 2)
    if abs(v) >= 1000:
        return f"{v:,.0f}"
    return _trim(v, 4)


def sig2(x: float) -> float:
    """Round to two significant digits: 0.006375 -> 0.0064, 2604971 -> 2600000."""
    if not x or not math.isfinite(x):
        return 0.0 if not x else x
    return round(x, 1 - int(math.floor(math.log10(abs(x)))))


# --------------------------------------------------------------------------- #
# the compiled workbook
# --------------------------------------------------------------------------- #
class ModelGraph:
    """One pycel compiler per workbook, with a lock and a promise to put values back.

    `compile()` does three things: build the dependency graph over every formula cell,
    recompute each formula once to find the cells pycel cannot evaluate (those are what
    `partial` reports), and snapshot every value so a what-if can always be undone.
    """

    def __init__(self, path: Path | str, file: str = "", key: str = ""):
        self.path = Path(path)
        self.file = file or self.path.name
        self.key = key                       # content hash, so a changed file rebuilds
        self.error: str | None = None
        self.compile_ms = 0.0
        self.broken: dict[str, str] = {}     # address -> why it cannot be evaluated
        self.mismatched: dict[str, tuple] = {}   # address -> (excel, pycel)
        self.lock = threading.RLock()
        self._ex = None
        self._snapshot: dict = {}
        self._tainted: set[str] = set()
        self._thread: threading.Thread | None = None
        self._context: tuple | None = None

    # ---------------------------------------------------------------- lifetime
    @property
    def ready(self) -> bool:
        return self._ex is not None

    def start(self) -> None:
        """Compile in the background. Safe to call repeatedly; never raises."""
        with self.lock:
            if self._ex is not None or (self._thread and self._thread.is_alive()):
                return
            self._thread = threading.Thread(target=self._safe_compile, daemon=True,
                                            name=f"scratch:{self.file}")
            self._thread.start()

    def _safe_compile(self) -> None:
        try:
            self.compile()
        except Exception as e:                     # a bad workbook must not kill the app
            log.warning("scratch: %s did not compile: %s", self.file, e)

    def compile(self) -> "ModelGraph":
        with self.lock:
            if self._ex is not None:
                return self
            t0 = time.perf_counter()
            try:
                self._compile()
            except Exception as e:
                self.error = f"{type(e).__name__}: {e}"
                self._ex = None
                raise
            self.compile_ms = (time.perf_counter() - t0) * 1000
            return self

    def _compile(self) -> None:
        import networkx as nx
        from pycel import ExcelCompiler
        from pycel.excelcompiler import _Cell

        # pycel logs a traceback for every formula that raises; we count them instead.
        logging.getLogger("pycel").setLevel(logging.CRITICAL)
        ex = ExcelCompiler(str(self.path), plugins=[PLUGIN])

        for addr in ex.formula_cells():
            try:
                ex.evaluate(addr.address)
            except Exception as e:                 # recorded again in the pass below
                self.broken[addr.address] = f"{type(e).__name__}: {_one_line(e)}"

        # Every formula cell recomputed once from its (cached) precedents. A cell that
        # raises here is one pycel cannot run at all: anything under it in a what-if is
        # still carrying Excel's last saved number, which is what `partial` warns about.
        for cell in [c for c in ex.cell_map.values()
                     if isinstance(c, _Cell) and c.python_code
                     and not c.address.is_unbounded_range]:
            addr = cell.address.address
            original = cell.value
            cell.value = None
            try:
                ex.evaluate(addr)
                if original is not None and not cell.close_enough(original):
                    self.mismatched[addr] = (original, cell.value)
            except Exception as e:
                self.broken[addr] = f"{type(e).__name__}: {_one_line(e)}"
            finally:
                cell.value = original

        self._ex = ex
        self._snapshot = {k: c.value for k, c in ex.cell_map.items()}
        self._tainted = set(self.broken)
        for addr in self.broken:
            node = ex.cell_map.get(addr)
            if node is not None:
                self._tainted |= {c.address.address for c in nx.descendants(ex.dep_graph, node)}

    # ------------------------------------------------------------------ access
    def context(self, con, path: Path | str, file: str) -> tuple:
        """(the `cells` rows, the number format of each) for this workbook, built once."""
        if self._context is None:
            book = Sheet(con, file)
            self._context = (book, number_formats(Path(path), set(book.by_ref)))
        return self._context

    def _node(self, sheet: str, cell: str):
        return self._ex.cell_map.get(canon(sheet, cell))

    def known(self, sheet: str, cell: str) -> bool:
        return self.ready and self._node(sheet, cell) is not None

    def is_formula(self, sheet: str, cell: str) -> bool:
        node = self._node(sheet, cell)
        return bool(node is not None and getattr(node, "formula", None))

    def value(self, sheet: str, cell: str):
        node = self._node(sheet, cell)
        return None if node is None else node.value

    def unstable(self, sheet: str, cell: str) -> bool:
        """pycel recomputes this cell to something other than Excel's saved value."""
        return canon(sheet, cell) in self.mismatched

    # ------------------------------------------------------------------ what-if
    def whatif(self, sets: dict[str, float], watch: list[str]) -> dict:
        """{ref: new value} for `watch` after applying `sets`. Always restores.

        `sets` and `watch` are canonical `Sheet!A1` strings. Cells that cannot be
        evaluated come back with their cached value and are listed in `partial`.
        """
        if not self.ready:
            self.compile()
        with self.lock:
            t0 = time.perf_counter()
            values: dict[str, object] = {}
            partial: list[str] = []
            try:
                for ref, val in sets.items():
                    self._ex.set_value(ref, val)
                self._repair()
                for ref in watch:
                    try:
                        values[ref] = self._ex.evaluate(ref)
                    except Exception as e:
                        values[ref] = self._snapshot.get(ref)
                        partial.append(ref)
                        log.debug("scratch: %s could not be recalculated: %s", ref, e)
            finally:
                self._restore()
            for ref in watch:
                if ref in self._tainted and ref not in partial:
                    partial.append(ref)
            return {"values": values, "partial": partial,
                    "ms": (time.perf_counter() - t0) * 1000}

    def _repair(self) -> None:
        """Put the cached value back under every cell pycel cannot run.

        `set_value` blanks the whole downstream cone. A blanked cell that raises on
        recalculation would take the entire evaluation with it, so the unrunnable ones go
        back to the workbook's number and everything under them is reported as `partial`.
        """
        for addr in self.broken:
            node = self._ex.cell_map.get(addr)
            if node is not None and node.value is None:
                node.value = self._snapshot.get(addr)

    def _restore(self) -> None:
        for addr, val in self._snapshot.items():
            node = self._ex.cell_map.get(addr)
            if node is not None:
                node.value = val
        for addr, node in self._ex.cell_map.items():       # cells born during the what-if
            if addr not in self._snapshot and getattr(node, "formula", None):
                node.value = None

    # ----------------------------------------------------------------- walking
    def downstream(self, sheet: str, cell: str, want,
                   keep_cap: int = BFS_KEEP_CAP) -> list[tuple[int, str]]:
        """Breadth-first over the dependency graph -> [(depth, address)], nearest first.

        `want(address)` picks the keepers. A whole level is always finished before
        stopping, so callers can group the siblings that sit at the same distance.
        """
        if not self.ready:
            return []
        start = self._node(sheet, cell)
        if start is None:
            return []
        seen, out, budget = {start}, [], BFS_NODE_BUDGET
        frontier, depth = [start], 0
        while frontier and len(out) < keep_cap and budget > 0:
            depth += 1
            nxt = []
            for node in frontier:
                for child in self._ex.dep_graph.successors(node):
                    budget -= 1
                    if child in seen or budget <= 0:
                        continue
                    seen.add(child)
                    nxt.append(child)
                    addr = child.address.address
                    if ":" not in addr and want(addr):
                        out.append((depth, addr))
            frontier = nxt
        return out

    def dependents(self, sheet: str, cell: str) -> int:
        import networkx as nx
        if not self.ready:
            return 0
        node = self._node(sheet, cell)
        if node is None:
            return 0
        try:
            return len(nx.descendants(self._ex.dep_graph, node))
        except Exception:
            return 0


def _formula_text(node) -> str:
    """The formula as Excel writes it, not as pycel compiles it."""
    formula = getattr(node, "formula", None)
    if formula is None:
        return ""
    return str(getattr(formula, "base_formula", None) or formula)[:160]


def _one_line(exc: Exception) -> str:
    return str(exc).strip().splitlines()[-1][:120] if str(exc).strip() else repr(exc)


# --------------------------------------------------------------------------- #
# lever discovery
# --------------------------------------------------------------------------- #
class Sheet:
    """Everything the `cells` table already knows about one workbook, in memory."""

    def __init__(self, con, file: str):
        self.rows = list(con.execute(
            "SELECT sheet, cell_ref, row, col, row_label, header, value_num, value_text, "
            "is_formula, formula FROM cells WHERE file=?", (file,)))
        self.by_ref = {(r["sheet"], r["cell_ref"]): r for r in self.rows}
        self.sheets: list[str] = []
        for r in self.rows:
            if r["sheet"] not in self.sheets:
                self.sheets.append(r["sheet"])

    def label(self, sheet: str, ref: str) -> str:
        r = self.by_ref.get((sheet, ref))
        return (r["row_label"] or "").strip() if r is not None else ""

    def row_text_left(self, sheet: str, row: int, col: int, left: int = 6) -> str:
        """The nearest text to the left of a cell on its own row.

        A scenario selector sits beside its name rather than under it (`A4 = "Market
        Scenario"`, `B4 = "Base"`), and that name is the only thing a client says.
        """
        cells = self._pos()
        for cc in range(col - 1, max(0, col - left) - 1, -1):
            r = cells.get((sheet, row, cc))
            if r is None:
                continue
            if r["value_num"] is not None or r["is_formula"]:
                return ""
            text = (r["value_text"] or "").strip()
            if text and len(text) <= 48:
                return text
            if text:
                return ""
        return ""

    def section(self, sheet: str, row: int, col: int, up: int = SECTION_UP) -> str:
        """The block heading above a cell: the first column-A line with no value beside it.

        `Structure 2: Percentage of GMV. Comps` is where the words a client actually says
        ("GMV fee") live, when the row label is the source's own wording. It is also the only
        thing that separates two cells holding the same 0.85% "platform fee" on two sheets,
        so it is kept on the lever itself and not only mined for aliases.

        Fifteen rows, not forty: a heading further up than that is another block's, and on
        the big model the longer reach changed three levers and every one of them for the worse.
        """
        cells = self._pos()
        for rr in range(row - 1, max(0, row - up) - 1, -1):
            r = cells.get((sheet, rr, 1))
            if r is None or r["is_formula"] or r["value_num"] is not None:
                continue
            text = (r["value_text"] or "").strip()
            if not text or len(text) > 70:
                continue
            beside = any(cells.get((sheet, rr, c)) is not None
                         and (cells[(sheet, rr, c)]["value_num"] is not None
                              or cells[(sheet, rr, c)]["value_text"])
                         for c in range(2, 6))
            if not beside:
                return text
        return ""

    def _pos(self) -> dict:
        cells = self.__dict__.get("_by_pos")
        if cells is None:
            cells = {(r["sheet"], r["row"], r["col"]): r for r in self.rows}
            self._by_pos = cells
        return cells

    def local_header(self, sheet: str, row: int, col: int, up: int = 15) -> str:
        """The nearest short text above a cell in its own column: the "Downside | Base | Upside"
        row of a scenario block, which the table pass misses when a sheet holds several blocks."""
        cells = self._pos()
        for rr in range(row - 1, max(0, row - up) - 1, -1):
            r = cells.get((sheet, rr, col))
            if r is None or r["value_num"] is not None or r["is_formula"]:
                continue                   # the block's other input rows sit between us and the header
            t = (r["value_text"] or "").strip()
            if t and len(t) <= 16 and len(t.split()) <= 2:
                return t
            if t:
                return ""                  # a sentence above is a title, not a column header
        return ""


def number_formats(path: Path, wanted: set[tuple[str, str]]) -> dict[tuple[str, str], str]:
    """One openpyxl pass for the number format of the cells we care about.

    `cells` does not store the format and this is the only thing that tells a rate from a
    dollar amount, so it is read here rather than widening the ingest schema.
    """
    import openpyxl

    out: dict[tuple[str, str], str] = {}
    by_sheet: dict[str, set[str]] = {}
    for sheet, ref in wanted:
        by_sheet.setdefault(sheet, set()).add(ref)
    wb = openpyxl.load_workbook(path, data_only=False)
    try:
        for sheet, refs in by_sheet.items():
            if sheet not in wb.sheetnames:
                continue
            ws = wb[sheet]
            for ref in refs:
                try:
                    cell = ws[ref]
                except Exception:
                    continue
                out[(sheet, ref)] = cell.number_format or "General"
    finally:
        wb.close()
    return out


def validation_lists(path: Path) -> dict[tuple[str, str], list[str]]:
    """Every dropdown in the workbook: {(sheet, 'B4'): ['Downside', 'Base', 'Upside']}.

    The dropdown is the author's own statement of what a selector may hold, so it beats
    anything inferred from the formulas. Inline lists only (`"a,b,c"`); a list that points
    at a range is left to the formula scan, which reads the same strings out of the IFs.
    """
    import openpyxl
    from openpyxl.utils.cell import range_boundaries

    out: dict[tuple[str, str], list[str]] = {}
    try:
        wb = openpyxl.load_workbook(path, data_only=False)
    except Exception as e:  # noqa: BLE001
        log.debug("scratch: no validations from %s: %s", path, e)
        return out
    try:
        for ws in wb.worksheets:
            for dv in list(getattr(ws.data_validations, "dataValidation", []) or []):
                if str(getattr(dv, "type", "")) != "list":
                    continue
                raw = str(dv.formula1 or "").strip()
                if not (raw.startswith('"') and raw.endswith('"')):
                    continue
                choices = [c.strip() for c in raw[1:-1].split(",") if c.strip()]
                if len(choices) < 2:
                    continue
                for rng in getattr(dv.sqref, "ranges", []) or []:
                    lo_c, lo_r, hi_c, hi_r = range_boundaries(str(rng.coord))
                    if (hi_c - lo_c + 1) * (hi_r - lo_r + 1) > 64:
                        continue
                    for c in range(lo_c, hi_c + 1):
                        for r in range(lo_r, hi_r + 1):
                            out[(ws.title, f"{col_letters(c)}{r}")] = choices[:MAX_CHOICES]
    finally:
        wb.close()
    return out


_MIRROR = re.compile(r"^=\s*(?:'[^']+'|[A-Za-z_][A-Za-z0-9_.]*)!\$?[A-Z]{1,3}\$?\d+\s*$"
                     r"|^=\s*\$?[A-Z]{1,3}\$?\d+\s*$")


def mirror_groups(book: "Sheet") -> dict[tuple[str, str], set[tuple[str, str]]]:
    """Cells that are nothing but a copy of another cell, grouped with their source.

    A selector on the summary sheet is echoed onto every model tab (`Assumptions!B3 =
    'Summary Scenarios'!$B$4`) and the IFs compare the echo, not the original. Without the
    group, the original looks like a cell nothing reads.
    """
    parent: dict[tuple[str, str], tuple[str, str]] = {}
    for r in book.rows:
        if not r["is_formula"] or not r["formula"]:
            continue
        if not _MIRROR.match(str(r["formula"]).strip()):
            continue
        refs = refs_in_formula(r["formula"], r["sheet"])
        if len(refs) != 1:
            continue
        src = next(iter(refs))
        if src in book.by_ref:
            parent[(r["sheet"], r["cell_ref"])] = src

    def root(key):
        seen = {key}
        while key in parent and parent[key] not in seen:
            key = parent[key]
            seen.add(key)
        return key

    groups: dict[tuple[str, str], set[tuple[str, str]]] = {}
    for child in parent:
        groups.setdefault(root(child), set()).add(child)
    out: dict[tuple[str, str], set[tuple[str, str]]] = {}
    for src, kids in groups.items():
        whole = kids | {src}
        for key in whole:
            out[key] = whole
    return out


_COMPARED = re.compile(r'(?:=|<>)\s*"([^"]{1,60})"|"([^"]{1,60})"\s*(?:=|<>)')


def compared_literals(readers: dict, group: set[tuple[str, str]]) -> list[str]:
    """The strings formulas test a selector against, in the order the sheet writes them."""
    out: list[str] = []
    seen: set[int] = set()
    for key in group:
        for r in readers.get(key, ()):  # noqa: B007
            if id(r) in seen:
                continue
            seen.add(id(r))
            for m in _COMPARED.finditer(r["formula"]):
                text = (m.group(1) or m.group(2) or "").strip()
                if text and text not in out and not text.endswith(" "):
                    out.append(text)
    return out[:MAX_CHOICES]


def match_choice(value, choices: list[str]) -> str | None:
    """"downside" / "flat dues" -> the workbook's own spelling, or None."""
    want = str(value or "").strip().lower()
    if not want:
        return None
    for c in choices:
        if c.lower() == want:
            return c
    starts = [c for c in choices if c.lower().startswith(want)]
    if len(starts) == 1:
        return starts[0]
    holds = [c for c in choices if want in c.lower()]
    return holds[0] if len(holds) == 1 else None


def _is_summary_sheet(name: str) -> bool:
    return bool(SUMMARY_SHEET_RE.search(name or ""))


def output_sheets(book: Sheet) -> set[str]:
    """Where a partner reads a number off: the summary-ish sheets, else the first with
    formulas."""
    named = {s for s in book.sheets if _is_summary_sheet(s)}
    if named:
        return named
    for sheet in book.sheets:
        if any(r["is_formula"] for r in book.rows if r["sheet"] == sheet):
            return {sheet}
    return set()


def pick_outputs(hits: list[tuple[int, str]], book: Sheet, limit: int) -> list[str]:
    """One cell per metric: the row a partner reads it off, and the column they quote.

    A lever lands on a whole row at once (`Subscriber fees` across Year 0 to Year 5) and often
    on the same row several times over, because a summary sheet restates each metric under
    every option it prices: a selected-scenario block at the top, then one block per fee
    structure below. Keeping all of them spends the eight slots on one metric, so:

      one cell per label.  Among the rows carrying that label, the highest one wins: a
      summary sheet states the selected case first and the alternatives underneath, so the
      top row is the number on the page the client is looking at. (Before this rule the
      nearest row in the dependency graph won, which picked the per-structure block and
      answered a question about the recommended case with another case's figure.)

      one column per row.  The middle when the row is narrow (three columns is almost
      always low / base / high) and the last when it is wide (a time series, where the
      terminal period is the number people quote).

    Metrics come back nearest-first in the dependency graph, which puts what the lever
    drives directly above what it drives through five other cells.
    """
    groups: dict[str, dict[int, list[tuple[int, str, int]]]] = {}
    for depth, addr in hits:
        sheet, ref = split_ref(addr)
        row = book.by_ref.get((sheet, ref))
        if row is None:
            continue
        label = book.label(sheet, ref)
        groups.setdefault(label, {}).setdefault(row["row"], []).append(
            (row["col"], addr, depth))

    chosen: list[tuple[int, int, int, str]] = []
    for label, rows in groups.items():
        top = min(rows)
        subscribers = sorted(rows[top])
        col, addr, _ = subscribers[len(subscribers) // 2] if len(subscribers) <= 3 else subscribers[-1]
        chosen.append((min(d for _, _, d in subscribers), top, col, addr))
    chosen.sort()
    return [addr for _, _, _, addr in chosen[:limit]]


def _drop_echoes(outs: list[dict], label: str, value: float) -> list[dict]:
    """A summary sheet that simply restates the input is not an output.

    `Summary Recommendation!B72 = Assumptions!B49` would otherwise make the lever answer
    its own question ("addressable share goes from 45% to 34%"), so a cell carrying the
    lever's own label or its exact value drops to the back of the list.
    """
    real = [o for o in outs
            if o["label"].strip().lower() != label.strip().lower()
            and not (isinstance(o["value"], (int, float))
                     and math.isclose(float(o["value"]), value, rel_tol=1e-12))]
    return real + [o for o in outs if o not in real]


# --------------------------------------------------------------------------- #
# spoken names, from the workbook's own words
# --------------------------------------------------------------------------- #
_QUALIFIER = {"median", "blended", "average", "annual", "ongoing", "combined", "total",
              "loaded", "target", "selected", "estimated", "assumed", "effective",
              "comparable", "current", "gross", "net", "first", "core", "overall"}
_CUT_AT = re.compile(r"\s+(?:on|per|over|for|from|as|of|through|during|under|by|in|at|to|"
                     r"paying|spread|including)\s+|,|\(|:|–|—|-\s", re.I)
_ACRONYM = re.compile(r"\b([A-Z]{2,6})\b")
_PARENS = re.compile(r"\(([^)]{2,24})\)")
ALIAS_STOP = {"the", "a", "an", "and", "or", "of", "to", "in", "on", "for", "at", "by",
              "with", "as", "is", "are", "per", "from", "that", "this", "it", "its", "one"}


def _alias_ok(name: str, label: str = "", out: list[str] | None = None,
              least: int = 1) -> str:
    """One candidate spoken name, cleaned, or "" when it is not one.

    Up to three words, at least `least` of them, at least one carrying meaning, never the
    label again, and never a bare index: a digit in a block title ("Structure 3: Year 5 at
    Selected Levers") numbers the block, and "year 5" as a name answers every line with a 5
    in it.
    """
    name = re.sub(r"\s+", " ", str(name or "")).strip(" .,:;-").lower()
    words = name.split()
    if not (least <= len(words) <= 3) or not any(w not in ALIAS_STOP for w in words):
        return ""
    if any(w.isdigit() for w in words):
        return ""
    if len(name) < 3 or name == (label or "").strip().lower() or name in (out or []):
        return ""
    return name


_BLOCK_NOUN = re.compile(r"\b(comps?|inputs?|assumptions?|levers?|parameters?|notes?|detail|"
                         r"details|summary|outputs?|calcs?|calculations?|data|table)\b", re.I)


def section_aliases(section: str) -> list[str]:
    """The content words of a block title, as one or two names a client would say.

    "Structure 2: Percentage of GMV. Comps" is how a spreadsheet says what a client calls
    "percentage of GMV", and the row label under it ("Median platform fee on portfolios over
    $1M") never says it. The index before the colon is dropped, the boilerplate noun after
    it is dropped, and an "and" splits a heading that names two things.

    Two words at least. One word out of a block title names the block and not the row, and
    every row under it inherits it: "Build-Up: Design and UX" gave "design" to the design
    *hours* and it outranked the design *rate* on a line about the rate (E36).
    """
    out: list[str] = []
    text = str(section or "").split(":", 1)[-1]
    for part in re.split(r"[.;,]|\band\b", text):
        part = _BLOCK_NOUN.sub(" ", part).strip(" .,:;-")
        if not part:
            continue
        name = _alias_ok(part, out=out, least=2) or _alias_ok(
            _CUT_AT.split(part, maxsplit=1)[0], out=out, least=2)
        if name:
            out.append(name)
        if len(out) >= 2:
            break
    return out


def context_of(sheet: str, section: str) -> str:
    """Where a lever lives, in one string: the sheet and the block inside it.

    `Subscriber Fee Model › Structure 2: Percentage of GMV. Comps` against
    `Subscriber Ramp › Subscriber Volume Inputs`. Two cells labelled "platform fee" and
    holding 0.85% are the same lever to anything that reads only the label.
    """
    sheet = str(sheet or "").strip()
    section = str(section or "").strip()
    if section and section.lower() != sheet.lower():
        return f"{sheet} › {section}"
    return sheet


def local_aliases(label: str, section: str = "") -> list[str]:
    """Two or three names a client would say, taken out of the sheet, with no model call.

    The row label is the source's wording ("Median platform fee on accounts over $1M");
    what gets said on the call is the head of it ("platform fee") and the acronym in the
    block heading above it ("Structure 2: Percentage of GMV" -> "GMV"). Both are in the
    file, so the lever table matches spoken English before the deep pass has run, and
    identically on every machine. The model call still runs and adds to these.
    """
    out: list[str] = []

    def add(name: str) -> None:
        name = re.sub(r"\s+", " ", str(name or "")).strip(" .,:;-").lower()
        words = name.split()
        if not (1 <= len(words) <= 3) or not any(w not in ALIAS_STOP for w in words):
            return
        if len(name) < 3 or name == (label or "").strip().lower() or name in out:
            return
        out.append(name)

    for m in _PARENS.finditer(label or ""):                  # "... (CAC)"
        add(m.group(1))
    head = _CUT_AT.split(str(label or ""), maxsplit=1)[0].strip()
    add(head)
    words = head.split()
    if len(words) > 2 and words[0].lower() in _QUALIFIER:    # "Median platform fee"
        add(" ".join(words[1:]))
    # Acronyms only out of the block heading: the rest of a heading is a sentence about the
    # block ("Structure 2: Percentage of GMV"), and a phrase out of it matches half the
    # sheet, while the acronym in it is exactly what gets said.
    for m in _ACRONYM.finditer(f"{str(section or '').split(':')[-1]} {label or ''}"):
        add(m.group(1))
    return out[:3]


def _fold(word: str) -> str:
    """Plurals folded the way the matcher folds them, so "years" and "year" count once."""
    word = word.strip(".").lower()
    return word[:-1] if len(word) > 3 and word.endswith("s") and not word.endswith("ss") else word


_SELECTOR_NOUN = re.compile(r"\b(setting|scenario|model|positioning|structure|entity|mode|"
                            r"case|option|choice|type|switch|selector|method|basis)\b", re.I)
_YES_NO = {"yes", "no", "y", "n", "on", "off", "true", "false", "none"}


def _selector_label(label: str, choices: list[str]) -> str:
    """Say that a two-word selector is a setting, because otherwise it answers to its words.

    "Subscriber Fee" holds one of four named fee structures, and a client asking what happens
    to subscriber fees at a different rate hit the selector instead of the rate: two words, both
    of them said, a perfect match. One more word breaks the tie and is true of the cell. A
    yes/no toggle keeps its own label, where the extra word would only mislead.
    """
    words = str(label or "").split()
    if len(words) > 2 or _SELECTOR_NOUN.search(label or ""):
        return label
    if all(str(c).strip().lower() in _YES_NO for c in choices):
        return label
    return f"{label} setting"


def _choice_alias(choice: str) -> str:
    """"Channel + Union Split (S2)" -> "channel union split": what a client would say of it."""
    text = re.sub(r"\(\s*S\d+\s*\)", " ", str(choice or ""))
    text = re.sub(r"[^A-Za-z0-9 +-]", " ", text).replace("+", " ")
    words = [w for w in text.lower().split() if w not in ALIAS_STOP]
    return " ".join(words[:3])


def _prune_aliases(levers: list[dict]) -> None:
    """Drop a name that half the table would answer to ("annual fee", "years").

    A spoken name only earns its place if it separates one lever from the rest, and a name
    is scored against the whole utterance, so a one-word name that half the sheet also uses
    matches everything: "Years to exit" shortened to "years" answered every question with
    the word year in it and outranked the lever the client meant. One word therefore has to
    be near-unique in the table; two or three words only need one uncommon word between
    them. Measured on the big model: with this rule the table's top pick is right 20 points more often
    than with unrestricted short names.
    """
    # A block title is shared by every row under it by construction, so counting its words here
    # would say they identify nothing and prune the names that do: on the big model the fifteen rows under
    # "Channel and Market Data" pushed "market" over the ceiling and took it off the pricing
    # switch, whose two values are Premium and Market (E36). Block names are counted out of the
    # frequencies; they are still pruned by them.
    block = {id(lv): set(section_aliases(lv.get("section") or "")) for lv in levers}
    df: dict[str, int] = {}
    for lv in levers:
        names = [lv.get("label") or ""] + [a for a in (lv.get("aliases") or [])
                                           if a not in block[id(lv)]]
        for word in {_fold(w) for n in names for w in re.findall(r"[a-z0-9%.]+", n.lower())}:
            df[word] = df.get(word, 0) + 1
    ceiling = max(3, len(levers) // 8)
    for lv in levers:
        keep = []
        for alias in lv.get("aliases") or []:
            words = [_fold(w) for w in alias.split()]
            limit = max(3, len(levers) // 32) if len(words) == 1 else ceiling
            if any(df.get(w, 0) <= limit for w in words):
                keep.append(alias)
        # a selector's names are its own plus one per value it may hold, and the values are
        # what a client actually says, so they must not be squeezed out by the section
        lv["aliases"] = keep[:MAX_ALIASES + MAX_CHOICES if lv.get("fmt") == "text"
                             else MAX_ALIASES]


def _rank_key(lever: dict) -> tuple:
    return (0 if lever["outputs"] else 1,
            0 if INPUT_SHEET_RE.search(lever["sheet"]) else 1,
            -lever["reach"], -lever["reads"], lever["cell"])


def _round_robin(levers: list[dict]) -> list[dict]:
    by_sheet: dict[str, list[dict]] = {}
    for lever in sorted(levers, key=_rank_key):
        by_sheet.setdefault(lever["sheet"], []).append(lever)
    order = sorted(by_sheet, key=lambda s: _rank_key(by_sheet[s][0]))
    out: list[dict] = []
    while len(out) < len(levers):
        for sheet in order:
            if by_sheet[sheet]:
                out.append(by_sheet[sheet].pop(0))
    return out


def _interleave(levers: list[dict]) -> list[dict]:
    """Levers that move a summary number first, then round-robin across sheets, so one
    roster table cannot fill the prompt block while the assumptions sheet waits."""
    return (_round_robin([lv for lv in levers if lv["outputs"]])
            + _round_robin([lv for lv in levers if not lv["outputs"]]))


def _picker(book: Sheet, graph: ModelGraph, targets: set[str]):
    """What counts as an output: labelled, numeric, a formula, on a summary-ish sheet, and
    a cell pycel reproduces (one it recomputes differently would report a phantom move)."""
    def is_output(addr: str) -> bool:
        try:
            sheet, ref = split_ref(addr)
        except ValueError:
            return False
        value = graph.value(sheet, ref)
        return (sheet in targets and bool(book.label(sheet, ref))
                and isinstance(value, (int, float)) and not isinstance(value, bool)
                and graph.is_formula(sheet, ref) and not graph.unstable(sheet, ref))
    return is_output


def outputs_for(con, path: Path, file: str, graph: ModelGraph, sheet: str, cell: str,
                limit: int = MAX_OUTPUTS) -> list[dict]:
    """The labelled summary cells one input moves, nearest first."""
    book, formats = graph.context(con, path, file)
    targets = output_sheets(book)
    outs = []
    for addr in pick_outputs(graph.downstream(sheet, cell, _picker(book, graph, targets)),
                             book, limit):
        osheet, oref = split_ref(addr)
        outs.append({"cell": canon(osheet, oref),
                     "label": book.label(osheet, oref),
                     "fmt": fmt_kind(formats.get((osheet, oref), ""),
                                     graph.value(osheet, oref)),
                     "value": graph.value(osheet, oref)})
    row = book.by_ref.get((sheet, cell))
    value = row["value_num"] if row is not None and row["value_num"] is not None else 0.0
    return _drop_echoes(outs, book.label(sheet, cell), float(value))


def discover(con, path: Path, file: str, graph: ModelGraph) -> list[dict]:
    """The lever table for one workbook. CPU only; the graph must already be compiled."""
    # One openpyxl pass for every cell we might print: `cells` does not carry the number
    # format and it is the only thing that tells a rate from a dollar amount.
    book, formats = graph.context(con, path, file)
    refcount: dict[tuple[str, str], int] = {}
    readers: dict[tuple[str, str], list] = {}
    for r in book.rows:
        if not r["is_formula"] or not r["formula"]:
            continue
        for key in refs_in_formula(r["formula"], r["sheet"]):
            if key in book.by_ref:
                refcount[key] = refcount.get(key, 0) + 1
                readers.setdefault(key, []).append(r)

    candidates = []
    for r in book.rows:
        key = (r["sheet"], r["cell_ref"])
        if r["is_formula"] or r["value_num"] is None or r["value_text"] is not None:
            continue
        if not (r["row_label"] or "").strip() or not refcount.get(key):
            continue
        if not graph.known(*key) or graph.is_formula(*key):
            continue                 # array formulas look like plain numbers in `cells`
        candidates.append(r)

    levers = []
    per_label: dict[tuple[str, str], int] = {}
    for r in candidates:
        sheet, ref = r["sheet"], r["cell_ref"]
        label = r["row_label"].strip()
        seen = per_label.get((sheet, label), 0)
        if seen >= MAX_PER_LABEL:
            continue
        per_label[(sheet, label)] = seen + 1
        outs = outputs_for(con, path, file, graph, sheet, ref)
        header = (r["header"] or "").strip()
        if not header or header.lower().startswith("col ") or len(header) > 16 or len(header.split()) > 2:
            header = book.local_header(sheet, int(r["row"]), int(r["col"])) or header
        section = book.section(sheet, int(r["row"]), int(r["col"]))
        levers.append({
            "file": file, "sheet": sheet, "cell": ref, "label": label,
            "header": header, "section": section, "context": context_of(sheet, section),
            "value": float(r["value_num"]),
            "fmt": fmt_kind(formats.get((sheet, ref), ""), r["value_num"]),
            "choices": [],
            "aliases": local_aliases(label, section) + section_aliases(section),
            "outputs": outs,
            "reads": refcount[(sheet, ref)], "reach": graph.dependents(sheet, ref),
        })

    levers += discover_text(con, path, file, graph, refcount, readers)
    _prune_aliases(levers)
    return _interleave(levers)[:MAX_LEVERS]


def discover_text(con, path: Path, file: str, graph: ModelGraph,
                  refcount: dict[tuple[str, str], int], readers: dict) -> list[dict]:
    """Scenario selectors: the text cells that switch the model between named cases.

    `Summary Scenarios!B5 = "Premium"` is read by 39 formulas and swings year 5 subscriber
    revenue from $8.1M to $2.9M, which no numeric lever in the table can do. What makes it
    a lever rather than a caption is that the workbook says what it may hold: a dropdown
    list on the cell, or, failing that, the literal strings the IFs compare it against.
    """
    book, _ = graph.context(con, path, file)
    lists = validation_lists(path)
    mirrors = mirror_groups(book)
    levers: list[dict] = []
    seen_groups: list[set[tuple[str, str]]] = []
    for r in book.rows:
        key = (r["sheet"], r["cell_ref"])
        if r["is_formula"] or r["value_num"] is not None or not (r["value_text"] or "").strip():
            continue
        if not graph.known(*key) or graph.is_formula(*key):
            continue
        group = mirrors.get(key, {key})
        if not sum(refcount.get(k, 0) for k in group):
            continue                        # a caption, not a switch
        if any(group & g for g in seen_groups):
            continue                        # the echo of a selector is not a second lever
        current = str(r["value_text"]).strip()
        dropdown = lists.get(key)
        literals = [c for c in compared_literals(readers, group) if c and len(c) <= 60]
        if not dropdown:
            # No dropdown: the formulas have to recognise the value that is in the cell,
            # or it is a caption that happens to sit next to a test. A yes/no toggle is the
            # exception, because only one side of it is ever written in the IF.
            yes_no = (current.lower() in _YES_NO
                      and any(c.lower() in _YES_NO for c in literals))
            if current not in literals and not yes_no:
                continue
        # the workbook's own order, which is the order they are read out on a call
        choices = list(dict.fromkeys(list(dropdown or literals) + [current]))[:MAX_CHOICES]
        if len(choices) < 2:
            continue
        outs = outputs_for(con, path, file, graph, r["sheet"], r["cell_ref"])
        if not outs:
            continue        # a caption in a reference table, not a switch anyone can pull
        seen_groups.append(group)
        row, col = int(r["row"]), int(r["col"])
        label = (r["row_label"] or "").strip() or book.row_text_left(r["sheet"], row, col)
        if not label:
            continue
        label = _selector_label(label, choices)
        header = (r["header"] or "").strip()
        if not header or header.lower().startswith("col ") or len(header) > 16:
            header = book.local_header(r["sheet"], row, col) or ""
        section = book.section(r["sheet"], row, col)
        levers.append({
            "file": file, "sheet": r["sheet"], "cell": r["cell_ref"], "label": label,
            "header": header, "section": section,
            "context": context_of(r["sheet"], section),
            "value": current, "fmt": "text", "choices": choices,
            # the names a client says for a selector are usually its values, not its own
            # name: "run the downside", "go to market pricing", "no channel fee". They go in
            # as spoken names so the matcher finds the cell from the case the client asked
            # for; _prune_aliases drops the ones too common to identify anything.
            "aliases": local_aliases(label, section) + section_aliases(section)
                       + [_choice_alias(c) for c in choices],
            "outputs": outs,
            "reads": sum(refcount.get(k, 0) for k in group),
            "reach": graph.dependents(r["sheet"], r["cell_ref"]),
        })
    return levers


# --------------------------------------------------------------------------- #
# the sensitivity grid
# --------------------------------------------------------------------------- #
def grid_points(value: float, kind: str) -> list[float]:
    """The handful of values worth precomputing for one lever: the round numbers a client would
    say first (1%, 1.25%, 10%, 12%), then plain multiples of the current value."""
    if not value or not math.isfinite(value):
        return []
    lo, hi = sorted((value / 2.0, value * 2.0))
    ranked: list[tuple[int, float]] = []          # (priority, value): 0 = a round spoken value
    if kind == "pct":
        for decade in range(-6, 2):                     # 1-2-5 style round rates
            for mult in (1, 2, 5):
                rate = mult * (10.0 ** decade)
                if lo <= rate <= hi:
                    ranked.append((0, rate))
        # quarter points below 5%, whole points to 20%, five-point steps above
        step = 0.0025 if value < 0.05 else (0.01 if value < 0.2 else 0.05)
        q = math.ceil(lo / step)
        while q * step <= hi:
            ranked.append((0, q * step))
            q += 1
        ranked += [(1, value * m) for m in (0.5, 0.75, 1.25, 1.5, 2.0)]
    else:
        ranked += [(1, value * m) for m in (0.75, 0.9, 1.1, 1.25, 1.5)]

    def sig3(v: float) -> float:
        return float(f"{v:.3g}")

    seen, keep = set(), []
    for pri, v in sorted(ranked, key=lambda t: (t[0], abs(t[1] - value))):   # round first, nearest first
        v = sig3(v)
        if v == 0 or math.isclose(v, value, rel_tol=1e-9) or v in seen:
            continue
        seen.add(v)
        keep.append(v)
    return sorted(keep[:MAX_GRID_PER_LEVER])


def grid_values(lever: dict) -> list:
    """The values worth precomputing for one lever: round numbers, or the other choices."""
    if lever.get("fmt") == "text":
        return [c for c in (lever.get("choices") or [])
                if str(c) != str(lever.get("value"))][:MAX_GRID_PER_LEVER]
    value = lever.get("value")
    if not isinstance(value, (int, float)) or isinstance(value, bool):
        return []
    return grid_points(float(value), lever.get("fmt") or "number")


def build_grid(graph: ModelGraph, levers: list[dict],
               budget: int = MAX_GRID_RECALCS) -> list[dict]:
    """Precompute each lever at a few round values. Ranked levers get the budget first."""
    grid: list[dict] = []
    for lever in levers:
        if budget <= 0:
            break
        if not lever["outputs"]:
            continue
        watch = [o["cell"] for o in lever["outputs"]]
        cell = canon(lever["sheet"], lever["cell"])
        for point in grid_values(lever):
            if budget <= 0:
                break
            budget -= 1
            try:
                got = graph.whatif({cell: point}, watch)
            except Exception as e:
                log.debug("scratch: grid point %s=%s failed: %s", cell, point, e)
                continue
            outs = {ref: val for ref, val in got["values"].items()
                    if isinstance(val, (int, float)) and not isinstance(val, bool)}
            if outs:
                grid.append({"cell": cell, "value": point, "outputs": outs,
                             "partial": got["partial"]})
    return grid


def build_combos(graph: ModelGraph, levers: list[dict], top: int = COMBO_LEVERS,
                 budget: int = MAX_COMBO_POINTS) -> list[dict]:
    """The top levers two at a time: "one percent and six percent churn" in one card.

    Clients move two things at once and expect one number back. A single-lever grid cannot
    answer that even in principle, so the headline levers are paired and each pair is run
    at its first round value on both sides. Pairs on the same sheet and row are skipped:
    Downside and Base of the same input are alternatives, not a combination.
    """
    picked = [lv for lv in levers
              if lv["outputs"] and lv.get("active") is not False and grid_values(lv)][:top]
    combos: list[dict] = []
    for i, a in enumerate(picked):
        for b in picked[i + 1:]:
            if budget <= 0:
                return combos
            if a["sheet"] == b["sheet"] and a["label"] == b["label"]:
                continue
            cells = {canon(a["sheet"], a["cell"]): grid_values(a)[0],
                     canon(b["sheet"], b["cell"]): grid_values(b)[0]}
            watch = list(dict.fromkeys([o["cell"] for o in a["outputs"]]
                                       + [o["cell"] for o in b["outputs"]]))[:MAX_OUTPUTS]
            budget -= 1
            try:
                got = graph.whatif(cells, watch)
            except Exception as e:  # noqa: BLE001
                log.debug("scratch: combo %s failed: %s", cells, e)
                continue
            outs = {ref: val for ref, val in got["values"].items()
                    if isinstance(val, (int, float)) and not isinstance(val, bool)}
            if outs:
                combos.append({"cells": cells, "outputs": outs, "partial": got["partial"]})
    return combos


# --------------------------------------------------------------------------- #
# goal seek: "what fee gets us to $10M"
# --------------------------------------------------------------------------- #
def bracket(value: float, lo: float | None, hi: float | None,
            wide: float = 8.0) -> tuple[float, float]:
    """A search range around the lever's own value when the asker did not give one.

    Scaled off the number already in the cell rather than guessed from a domain: a fee
    lives near 0.01 and a salary near 150,000, and the only thing they have in common is
    the current value. The range never crosses zero, because a model input that goes
    negative stops meaning anything (a discount rate of -100% took one solve to 1e72).
    """
    if lo is not None and hi is not None:
        return (float(lo), float(hi)) if lo <= hi else (float(hi), float(lo))
    if not value:
        span = 1.0 if lo is None and hi is None else abs(float(lo if lo is not None else hi)) or 1.0
        return (float(lo) if lo is not None else 0.0,
                float(hi) if hi is not None else wide * span)
    a, b = (value / wide, value * wide) if value > 0 else (value * wide, value / wide)
    return (float(lo) if lo is not None else a, float(hi) if hi is not None else b)


def solve(graph: ModelGraph, lever: str, target: str, want: float,
          lo: float | None = None, hi: float | None = None,
          tries: int = GOALSEEK_TRIES, tol: float = 1e-4) -> dict:
    """Bisection on the compiled graph: the lever value that puts `target` at `want`.

    Bisection rather than Newton because the model is full of thresholds, MINs and
    round-trips (headcount steps, a tax rate that switches at a floor): a derivative is
    meaningless across one of those, while a sign change is not. Each step is one full
    recalculation, about 3 ms, so sixty of them is well inside a call's patience.
    """
    t0 = time.perf_counter()
    seen: list[tuple[float, float]] = []

    def f(x: float) -> float | None:
        got = graph.whatif({lever: x}, [target])["values"].get(target)
        if isinstance(got, bool) or not isinstance(got, (int, float)):
            return None
        seen.append((x, float(got)))
        return float(got)

    current = graph.value(*split_ref(lever))
    current = float(current) if isinstance(current, (int, float)) and not isinstance(current, bool) else 0.0
    found = None
    for wide in (8.0, 64.0):
        a, b = bracket(current, lo, hi, wide)
        fa, fb = f(a), f(b)
        if fa is None or fb is None:
            return {"error": f"{target} is not a number under this what-if"}
        if (fa - want) * (fb - want) <= 0:
            found = (a, fa, b, fb)
            break
        # not bracketed by the ends: walk the range before widening, so a hump cannot hide it
        prev_x, prev_f = a, fa
        for i in range(1, 17):
            x = a + (b - a) * i / 16.0
            fx = f(x)
            if fx is None:
                continue
            if (prev_f - want) * (fx - want) <= 0:
                found = (prev_x, prev_f, x, fx)
                break
            prev_x, prev_f = x, fx
        if found or (lo is not None and hi is not None):
            break
    n = len(seen)
    if found is None:
        near = min(seen, key=lambda p: abs(p[1] - want))
        lo_s, hi_s = min(x for x, _ in seen), max(x for x, _ in seen)
        return {"error": f"{target} never reaches {want:,.2f} for {lever} between "
                         f"{lo_s:,.6g} and {hi_s:,.6g} (nearest {near[1]:,.2f} at {near[0]:,.6g})",
                "nearest": {"value": near[0], "achieved": near[1]},
                "tries": n, "ms": round((time.perf_counter() - t0) * 1000, 1)}
    a, fa, b, fb = found
    x = a
    for _ in range(tries):
        x = (a + b) / 2.0
        fx = f(x)
        n += 1
        if fx is None:
            return {"error": f"{target} stopped evaluating at {lever} = {x:,.6g}"}
        if abs(fx - want) <= max(tol * abs(want), 1e-9) or abs(b - a) <= 1e-12 * max(1.0, abs(x)):
            break
        if (fa - want) * (fx - want) <= 0:
            b, fb = x, fx
        else:
            a, fa = x, fx
    got = seen[-1][1] if seen else None
    monotone = all(y2 >= y1 for (_, y1), (_, y2) in zip(sorted(seen), sorted(seen)[1:])) or \
        all(y2 <= y1 for (_, y1), (_, y2) in zip(sorted(seen), sorted(seen)[1:]))
    return {"value": x, "achieved": got, "want": want, "monotone": bool(monotone),
            "tries": n, "ms": round((time.perf_counter() - t0) * 1000, 1)}


# --------------------------------------------------------------------------- #
# "why did it move": the formula chain from the lever to the number
# --------------------------------------------------------------------------- #
def _cell_preds(graph: ModelGraph, node) -> list:
    """The cells one cell reads, with pycel's range nodes stepped through.

    `=SUM(F63:F68)` records a single `F63:F68` predecessor; the cells are one level behind
    it. A walk that stops at the range stops at the first formula, which is no explanation.
    """
    out, seen = [], set()
    for p in graph._ex.dep_graph.predecessors(node):
        kids = ([k for k in graph._ex.dep_graph.predecessors(p)]
                if ":" in p.address.address else [p])
        for k in kids:
            addr = k.address.address
            if ":" in addr or addr in seen:
                continue
            seen.add(addr)
            out.append(k)
    return out


def trace(graph: ModelGraph, book: "Sheet", cell: str, sets: dict | None = None,
          hops: int = EXPLAIN_HOPS) -> list[dict]:
    """The shortest path through the dependency graph from a changed input to a cell.

    With `sets`, it is the route the change actually took, and each hop carries its value
    before and after, which is the whole answer to "why did costs move". Without `sets`, it
    is the cell's own precedents, deepest first: what the number is made of. Long chains
    are cut in the middle rather than the end, because the two cells nearest the answer are
    the ones a partner repeats.
    """
    import networkx as nx

    end = graph._node(*split_ref(cell))
    if end is None:
        return []
    path = []
    if sets:
        best = None
        for ref in sets:
            start = graph._node(*split_ref(ref))
            if start is None:
                continue
            try:
                route = nx.shortest_path(graph._ex.dep_graph, start, end)
            except Exception:                       # nx.NetworkXNoPath and friends
                continue
            if best is None or len(route) < len(best):
                best = route
        # a range node (`F63:F68`) is how pycel records "this formula reads a block"; it is
        # a step in the graph but not a cell anyone can be shown, so it drops out
        path = [n for n in (best or []) if ":" not in n.address.address]
    if not path:
        # what the number is made of: up the precedents, preferring another calculation over
        # an input, so the chain shows the mechanism rather than stopping at the first
        # constant it meets. Ties go to the precedent the rest of the model leans on most.
        node, path = end, [end]
        seen = {end}
        while len(path) < hops:
            preds = [p for p in _cell_preds(graph, node) if p not in seen
                     and isinstance(getattr(p, "value", None), (int, float))
                     and not isinstance(getattr(p, "value", None), bool)]
            if not preds:
                break
            formulas = [p for p in preds if getattr(p, "formula", None)]
            node = max(formulas or preds,
                       key=lambda p: graph.dependents(*split_ref(p.address.address)))
            seen.add(node)
            path.append(node)
        path.reverse()

    if len(path) > hops:                            # keep both ends, drop the middle
        keep = path[:max(1, hops // 2)] + path[-(hops - max(1, hops // 2)):]
        path = keep
    after = {}
    if sets:
        after = graph.whatif(dict(sets), [n.address.address for n in path])["values"]
    out = []
    for node in path:
        addr = node.address.address
        sheet, ref = split_ref(addr)
        was = graph.value(sheet, ref)
        now = after.get(addr, was)
        row = book.by_ref.get((sheet, ref)) if book else None
        label = (row["row_label"] or "") if row is not None else ""
        header = (row["header"] or "") if row is not None else ""
        delta = (now - was) if isinstance(now, (int, float)) and isinstance(was, (int, float)) \
            and not isinstance(now, bool) and not isinstance(was, bool) else None
        out.append({"cell": addr, "label": label, "header": header,
                    "formula": _formula_text(node), "was": was, "now": now, "delta": delta})
    return out


# --------------------------------------------------------------------------- #
# derived Q&A from the grid
# --------------------------------------------------------------------------- #
def _source_wording(label: str) -> bool:
    """True when a row label quotes its source rather than naming the input.

    "Median platform fee on accounts over $1M" and "Penn State and PNC banking
    partnership, annual royalty" are citations; "Annual subscriber churn" is a name.
    """
    label = str(label or "").strip()
    if not label:
        return False
    head = _CUT_AT.split(label, maxsplit=1)[0].strip()
    return head != label or len(label.split()) > 5


def _delta_word(now: float, was: float) -> str:
    if was is None or now is None or math.isclose(now, was, rel_tol=1e-9):
        return "holds at"
    return "rises to" if now > was else "falls to"


def _headline(outputs: list[dict], values: dict) -> tuple[dict | None, float]:
    """The output worth saying out of the ones a lever moves.

    A money or count figure that moves comes first (the nearest output can sit still when
    the model treats a downstream figure as a target); a ratio only when nothing else
    moves, because a margin swings on a small base; among equal movers the bigger figure
    is the headline (revenue over fee per subscriber).
    """
    best, best_rank = None, None
    for cand in outputs:
        v, w = values.get(cand["cell"]), cand.get("value")
        if not isinstance(v, (int, float)) or not isinstance(w, (int, float)):
            continue
        move = abs(v - w) / (abs(w) if w else 1.0)
        rank = (move >= 0.001, (cand.get("fmt") or "number") != "pct", round(move, 3), abs(w))
        if best_rank is None or rank > best_rank:
            best, best_rank = cand, rank
    return best, (best_rank[2] if best_rank else 0.0)


def derived_qa(record: dict) -> list[dict]:
    """Grid points -> Q&A rows in the same shape the deep cache already stores."""
    levers = {canon(lv["sheet"], lv["cell"]): lv for lv in record.get("levers", [])}
    siblings: dict[str, int] = {}
    for lv in record.get("levers", []):
        siblings[lv["label"]] = siblings.get(lv["label"], 0) + 1
    file = record.get("file", "")

    def spoken(lever: dict) -> str:
        head = str(lever.get("header") or "").strip()
        if siblings.get(lever["label"], 0) > 1 and head and len(head) <= 16 and len(head.split()) <= 2:
            return f"{lever['label']} ({head})"       # B39/C39/D39: Downside, Base, Upside
        return lever["label"]

    def asked(lever: dict) -> str:
        """The name the question is written under: the most specific one available.

        A comp row is labelled in the source's words ("Median platform fee on portfolios
        over $1M"), which is the one wording nobody says, so the block it sits under goes in
        beside it and the row reads "... (percentage of GMV)". That is what a line like
        "what if we charged 1% of GMV" has to land on.
        """
        name = spoken(lever)
        sec = next(iter(section_aliases(lever.get("section") or "")), "")
        if sec and _source_wording(lever.get("label")) and sec not in name.lower():
            return f"{name} ({sec})"
        return name

    out: list[dict] = []
    for point in record.get("grid", []):
        if point.get("cells"):
            out += _combo_qa(point, levers, spoken, file)
            continue
        lever = levers.get(point["cell"])
        if not lever or not lever["outputs"]:
            continue
        top, best_move = _headline(lever["outputs"], point["outputs"])
        if top is None:
            continue
        now = point["outputs"].get(top["cell"])
        was, kind = top.get("value"), top.get("fmt") or "number"
        new_in = fmt_value(point["value"], lever["fmt"])
        old_in = fmt_value(lever["value"], lever["fmt"])
        now_s, was_s = fmt_value(now, kind), fmt_value(was, kind)
        label = spoken(lever)
        # one row per spoken name: the label the sheet uses plus the aliases ("GMV fee",
        # "advisory rate"), so the way the client says it embeds close
        spoken_names = [a for a in (lever.get("aliases") or [])
                        if lever.get("fmt") != "text"
                        or not any(a == _choice_alias(c) for c in lever.get("choices") or [])]
        first = asked(lever)
        names = [first] + [a for a in spoken_names[:4]
                           if a and a.lower() not in (label.lower(), first.lower())]
        for i, name in enumerate(names):
            out.append({
                "question": (f"What if {name} were {new_in}?" if i == 0
                             else f"What if we changed the {name} to {new_in}?"),
                "answer": _words((f"{top['label']} goes from {was_s} to {now_s} when "
                                  f"{label} moves from {old_in} to {new_in}.") if best_move >= 0.001 else
                                 (f"Nothing on the summary moves: {top['label']} stays {was_s} when "
                                  f"{label} moves from {old_in} to {new_in}."), 25),
                "say": _words(f"{top['label']} {_delta_word(now, was)} {now_s}" if best_move >= 0.001
                              else f"{top['label']} stays {was_s}", 8),
                "figure": now_s,
                "citations": [{"file": file, "loc": pretty(lever["sheet"], lever["cell"])},
                              {"file": file, "loc": pretty(*split_ref(top["cell"]))}],
                "file": file,
                "loc": pretty(lever["sheet"], lever["cell"]),
                "derived": True,
                "lever_value": point["value"],
                "working": (f"{lever['sheet']}!{lever['cell']} {old_in} \u2192 {new_in}, "
                            "model recalculated"),
            })
    return out


def _combo_qa(point: dict, levers: dict, spoken, file: str) -> list[dict]:
    """A two-lever grid point as one Q&A row. `lever_value` stays None: no single number
    describes it, and the instant path must not serve it for one the client happened to say."""
    pair = [(levers.get(cell), value) for cell, value in point["cells"].items()]
    if any(lv is None or not lv["outputs"] for lv, _ in pair):
        return []
    outputs = list({o["cell"]: o for lv, _ in pair for o in lv["outputs"]}.values())
    top, best_move = _headline(outputs, point["outputs"])
    if top is None or best_move < 0.001:
        return []
    now, was = point["outputs"].get(top["cell"]), top.get("value")
    kind = top.get("fmt") or "number"
    parts = [f"{spoken(lv)} {fmt_value(value, lv['fmt'])}" for lv, value in pair]
    return [{
        "question": ("What if " + " and ".join(
            f"{spoken(lv)} were {fmt_value(value, lv['fmt'])}" for lv, value in pair) + "?"),
        "answer": _words(f"{top['label']} goes from {fmt_value(was, kind)} to "
                         f"{fmt_value(now, kind)} at {' and '.join(parts)}.", 25),
        "say": _words(f"{top['label']} {_delta_word(now, was)} {fmt_value(now, kind)}", 8),
        "figure": fmt_value(now, kind),
        "citations": [{"file": file, "loc": pretty(lv["sheet"], lv["cell"])} for lv, _ in pair]
                     + [{"file": file, "loc": pretty(*split_ref(top["cell"]))}],
        "file": file,
        "loc": pretty(pair[0][0]["sheet"], pair[0][0]["cell"]),
        "derived": True,
        "lever_value": None,
        "working": ", ".join(f"{lv['sheet']}!{lv['cell']} {fmt_value(lv['value'], lv['fmt'])} "
                             f"\u2192 {fmt_value(value, lv['fmt'])}" for lv, value in pair)
                   + ", model recalculated",
    }]


def _words(text: str, limit: int) -> str:
    parts = (text or "").split()
    return " ".join(parts[:limit])


# --------------------------------------------------------------------------- #
# the prompt block
# --------------------------------------------------------------------------- #
def levers_block(levers: list[dict], limit: int = 40) -> str:
    if not levers:
        return ""
    lines = ["== LEVERS (inputs whatif can change) =="]
    for lever in levers[:limit]:
        alias = ""
        names = [a for a in lever.get("aliases") or []
                 if lever.get("fmt") != "text"
                 or not any(a == _choice_alias(c) for c in lever.get("choices") or [])]
        if names:
            alias = " [" + "; ".join(names[:3]) + "]"
        value = fmt_value(lever["value"], lever["fmt"])
        if lever.get("fmt") == "text":
            value = f'"{value}"; one of ' + ", ".join(f'"{c}"' for c in lever["choices"][:8])
        where = f'({lever["file"]} \u203a {lever["sheet"]}!{lever["cell"]} = {value})'
        moves = ", ".join(f'{o["label"]} ({o["cell"]})' for o in lever["outputs"][:3])
        line = f'{lever["label"]}{alias} {where}'
        if moves:
            line += f" moves: {moves}"
        lines.append(line)
    return "\n".join(lines)


# --------------------------------------------------------------------------- #
# aliases: the one model call per workbook
# --------------------------------------------------------------------------- #
ALIASES_SCHEMA = {
    "type": "object",
    "properties": {
        "levers": {
            "type": "array",
            "items": {"type": "object",
                      "properties": {"i": {"type": "integer"},
                                     "aliases": {"type": "array",
                                                 "items": {"type": "string",
                                                           "maxLength": 40}}},
                      "required": ["i", "aliases"], "additionalProperties": False},
        }
    },
    "required": ["levers"],
    "additionalProperties": False,
}

# One call per workbook, and the reply has to fit inside classify()'s 1000-token output
# cap. Levers are numbered rather than named by cell: measured on Haiku, a cell string
# costs about eight output tokens and invites the model to echo the whole listing line
# back, which truncates the reply, and a truncated reply parses to nothing. Numbered, an
# entry costs about twenty tokens, so 32 levers land near 700.
ALIAS_BATCH = 32


def aliases_prompt(file: str, levers: list[dict]) -> str:
    listing = "\n".join(
        f'  [{i}] {lv["label"]} = {fmt_value(lv["value"], lv["fmt"])}  ({lv["sheet"]})'
        for i, lv in enumerate(levers[:ALIAS_BATCH]))
    return (
        f"These are the numbered input cells of {file} that a live what-if can change. For "
        "each, give 2 or 3 short names a client would SAY on a call when they mean that "
        "input.\n\n" + listing +
        "\n\nSpoken, not written: \"GMV fee\", \"advisory rate\", \"churn\", \"headcount "
        "ramp\". At most 3 words each, lower case, no cell references, no units, and no "
        "repeat of the label word for word. Return one entry per number above, as `i` (the "
        "number, nothing else) and `aliases`."
    )


def apply_aliases(levers: list[dict], raw: dict) -> int:
    """Merge a model reply into the lever table. A bad reply leaves the table usable.

    The names read off the sheet come first and are never dropped: they are the ones the
    matcher was measured on, and a model that returns nothing must not empty the table.
    Neither is a model name dropped. At a cap of four this merge quietly threw away the
    model's last name, and on the big model that was "gmv fee" on `Subscriber Fee Model!B48`, the one
    name that separated it from the other 0.85% platform fee in the workbook (E36).
    """
    n = 0
    for entry in (raw or {}).get("levers", []) or []:
        i = entry.get("i")
        if not isinstance(i, int) or not 0 <= i < min(len(levers), ALIAS_BATCH):
            continue
        names = [str(a).strip().lower()[:40] for a in (entry.get("aliases") or [])[:3]]
        names = [a for a in names if a and a != levers[i]["label"].lower()]
        if names:
            levers[i]["aliases"] = list(dict.fromkeys(
                (levers[i].get("aliases") or []) + names))[:MAX_ALIASES]
            n += 1
    return n


# --------------------------------------------------------------------------- #
# the whole scratch record for one workbook
# --------------------------------------------------------------------------- #
def _mark_active(levers: list[dict], grid: list[dict]) -> None:
    """`active` is False for an input a what-if cannot usefully move.

    Two ways to earn it: the summary does not read the input right now (the unselected
    price point, an unused scenario column), so every grid point moves nothing; or there is
    no value to try at all, which is what a hardcoded 0 in a year column is ("Subscriber fees,
    Year 0"). Both look like levers to a word matcher and neither can answer a question, so
    they are marked and the matcher's own penalty pushes them under the real one. A lever
    the grid ran out of budget for is left unmarked.
    """
    by_key = {canon(lv["sheet"], lv["cell"]): lv for lv in levers}
    seen: set[str] = set()
    moved: set[str] = set()
    for g in grid:
        if g.get("cells"):
            continue
        lv = by_key.get(g["cell"])
        if lv is None:
            continue
        seen.add(g["cell"])
        for o in lv.get("outputs") or []:
            v, w = g["outputs"].get(o["cell"]), o.get("value")
            if isinstance(v, (int, float)) and isinstance(w, (int, float)) and abs(v - w) > 1e-9 * max(1.0, abs(w)):
                moved.add(g["cell"])
                break
    for k, lv in by_key.items():
        if k in seen:
            lv["active"] = k in moved
        elif not grid_values(lv) or not lv.get("outputs"):
            lv["active"] = False


def _rank_outputs(levers: list[dict], grid: list[dict]) -> None:
    """Put the outputs a lever is measured to move in front of the ones it does not.

    The dependency graph says a cell is downstream; only a recalculation says it changes.
    Year 5 subscribers can be a target the model solves back to, so the nearest output of a
    churn lever sits still while the cost line moves.
    """
    moved: dict[str, set[str]] = {}
    for g in grid:
        if g.get("cells"):
            continue
        for ref, now in (g.get("outputs") or {}).items():
            lv_moved = moved.setdefault(g["cell"], set())
            if ref in lv_moved or not isinstance(now, (int, float)):
                continue
            lv_moved.add(ref)
    for lever in levers:
        key = canon(lever["sheet"], lever["cell"])
        base = {o["cell"]: o.get("value") for o in lever.get("outputs") or []}
        changed: set[str] = set()
        for g in grid:
            if g.get("cells") or g.get("cell") != key:
                continue
            for ref, now in (g.get("outputs") or {}).items():
                was = base.get(ref)
                if isinstance(now, (int, float)) and isinstance(was, (int, float)) \
                        and abs(now - was) > 1e-9 * max(1.0, abs(was)):
                    changed.add(ref)
        if changed:
            lever["outputs"] = ([o for o in lever["outputs"] if o["cell"] in changed]
                                + [o for o in lever["outputs"] if o["cell"] not in changed])


def build(con, path: Path, file: str, key: str = "",
          graph: ModelGraph | None = None) -> tuple[dict, ModelGraph]:
    """Compile, discover levers, run the grid. Pure CPU: call it in a thread."""
    t0 = time.perf_counter()
    graph = graph or ModelGraph(path, file, key)
    graph.compile()
    levers = discover(con, Path(path), file, graph)
    grid = build_grid(graph, levers)
    _mark_active(levers, grid)
    _rank_outputs(levers, grid)
    grid += build_combos(graph, _interleave(levers))
    record = {"file": file, "hash": key, "schema": SCHEMA, "levers": levers, "grid": grid,
              "built_ms": (time.perf_counter() - t0) * 1000,
              "compile_ms": round(graph.compile_ms, 1),
              "unevaluable": sorted(graph.broken)[:20],
              "mismatched": sorted(graph.mismatched)[:20]}
    return record, graph


# --------------------------------------------------------------------------- #
# calc: a grammar small enough to read in one sitting
# --------------------------------------------------------------------------- #
def _avg(*args):
    vals = _flat(args)
    if not vals:
        raise ValueError("calc: avg() needs at least one number")
    return sum(vals) / len(vals)


def _sum(*args):
    return sum(_flat(args))


def _cagr(start, end, years):
    if start == 0 or years == 0:
        raise ValueError("calc: cagr() needs a non-zero start and a non-zero number of years")
    if start < 0 or end < 0:
        raise ValueError("calc: cagr() needs positive values")
    return (end / start) ** (1.0 / years) - 1.0


def _pct(a, b):
    if b == 0:
        raise ValueError("calc: pct() cannot divide by zero")
    return a / b


def _flat(args):
    out = []
    for a in args:
        if isinstance(a, (list, tuple)):
            out.extend(_flat(a))
        else:
            out.append(float(a))
    return out


CALC_FUNCS = {
    "min": lambda *a: min(_flat(a)),
    "max": lambda *a: max(_flat(a)),
    "abs": lambda x: abs(float(x)),
    "round": lambda x, n=0: round(float(x), int(n)),
    "sum": _sum,
    "avg": _avg,
    "cagr": _cagr,          # cagr(start, end, years)
    "pct": _pct,            # pct(a, b) -> a as a share of b
}

_BINOPS = {ast.Add: lambda a, b: a + b, ast.Sub: lambda a, b: a - b,
           ast.Mult: lambda a, b: a * b, ast.Div: "div", ast.Pow: "pow", ast.Mod: "mod"}
MAX_EXPR_CHARS = 400


def eval_expr(expr: str, values: dict[str, float]) -> float:
    """Evaluate `expr` over `values`. Hand-walked AST; `eval` is never called.

    Numbers, names bound in `values`, `+ - * / ** %`, unary minus, parentheses and
    min/max/abs/round/sum/avg/cagr/pct. Anything else raises ValueError naming the token.
    """
    text = (expr or "").strip()
    if not text:
        raise ValueError("calc: the expression is empty")
    if len(text) > MAX_EXPR_CHARS:
        raise ValueError(f"calc: the expression is longer than {MAX_EXPR_CHARS} characters")
    try:
        tree = ast.parse(text, mode="eval")
    except SyntaxError as e:
        raise ValueError(f"calc: {text!r} is not an expression ({e.msg})") from None
    return float(_walk(tree.body, values))


def _walk(node, values: dict[str, float]):
    if isinstance(node, ast.Constant):
        if isinstance(node.value, bool) or not isinstance(node.value, (int, float)):
            raise ValueError(f"calc: {node.value!r} is not a number")
        return float(node.value)

    if isinstance(node, ast.Name):
        if node.id not in values:
            known = ", ".join(sorted(values)) or "none"
            raise ValueError(f"calc: '{node.id}' is not one of the inputs ({known})")
        return float(values[node.id])

    if isinstance(node, ast.UnaryOp):
        if isinstance(node.op, ast.USub):
            return -_walk(node.operand, values)
        if isinstance(node.op, ast.UAdd):
            return _walk(node.operand, values)
        raise ValueError(f"calc: unary '{type(node.op).__name__}' is not allowed")

    if isinstance(node, ast.BinOp):
        op = _BINOPS.get(type(node.op))
        if op is None:
            raise ValueError(f"calc: operator '{_op_text(node.op)}' is not allowed "
                             "(use + - * / ** %)")
        left, right = _walk(node.left, values), _walk(node.right, values)
        if op == "div":
            if right == 0:
                raise ValueError("calc: division by zero")
            return left / right
        if op == "mod":
            if right == 0:
                raise ValueError("calc: division by zero")
            return math.fmod(left, right)
        if op == "pow":
            if abs(right) > 100 or abs(left) > 1e12:
                raise ValueError("calc: that power is too large to evaluate")
            return left ** right
        return op(left, right)

    if isinstance(node, ast.Call):
        if not isinstance(node.func, ast.Name):
            raise ValueError(f"calc: '{_dotted(node.func)}' is not allowed")
        name = node.func.id
        if name not in CALC_FUNCS:
            allowed = ", ".join(sorted(CALC_FUNCS))
            raise ValueError(f"calc: '{name}' is not one of {allowed}")
        if node.keywords:
            raise ValueError(f"calc: {name}() does not take named arguments")
        args = [_walk(a, values) for a in node.args]
        try:
            return CALC_FUNCS[name](*args)
        except ValueError:
            raise
        except TypeError as e:
            raise ValueError(f"calc: {name}() got the wrong arguments ({e})") from None
        except ZeroDivisionError:
            raise ValueError("calc: division by zero") from None

    raise ValueError(f"calc: '{_node_text(node)}' is not allowed in an expression")


def _op_text(op) -> str:
    return {ast.BitOr: "|", ast.BitAnd: "&", ast.BitXor: "^", ast.FloorDiv: "//",
            ast.LShift: "<<", ast.RShift: ">>", ast.MatMult: "@"}.get(
        type(op), type(op).__name__)


def _dotted(node) -> str:
    try:
        return ast.unparse(node)
    except Exception:
        return type(node).__name__


def _node_text(node) -> str:
    try:
        return ast.unparse(node)[:60]
    except Exception:
        return type(node).__name__
