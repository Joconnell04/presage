"""Spreadsheets -> SQLite tables.

Consultant workbooks are messy, so nothing here assumes a clean rectangle:

  * merged cells are filled forward before anything is inspected
  * a sheet splits into blocks at 2+ blank rows and 2+ blank columns
  * inside a block the header is found by walking back from the first row that
    carries numbers, so title rows above the header are left out as `title`
  * a two-row header (a merged band over quarter labels) is joined into one label
  * formulas are read as cached values (data_only=True); a formula with no cached
    value is stored as NULL and counted

Every non-empty cell also lands in the universal long table `cells`, so a query
still works when block detection gets a sheet wrong.
"""

from __future__ import annotations

import csv
import datetime as dt
import re
from pathlib import Path

from .store import col_letter, ident, unique_ident

SHEET_EXT = {".xlsx", ".xlsm", ".xltx", ".xltm"}
CSV_EXT = {".csv", ".tsv"}
CELLS_MAX_DATA_ROWS = 5000        # do not blow up `cells` on a huge export
MAX_HEADER_ROWS = 2
# bump when what the ingest writes changes shape; every indexed file is rebuilt on the next open
INGEST_VERSION = 2                # 2: stacked tables split at their own header rows (2026-09-30)
TOTAL_RE = re.compile(r"^\s*(grand\s+)?(total|totals|sum|subtotal)\b|\btotals?\s*$", re.I)
_NUM_RE = re.compile(r"^\s*[\$£€]?\s*\(?-?[\d,]*\.?\d+\)?\s*%?\s*$")


def is_blank(v) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def as_number(v):
    """int/float straight through; money/percent/paren-negative strings parsed."""
    if isinstance(v, bool):
        return 1.0 if v else 0.0
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        s = v.strip()
        if not s or not _NUM_RE.match(s):
            return None
        neg = s.startswith("(") and s.endswith(")")
        pct = s.endswith("%")
        s = s.strip("()%").replace(",", "").replace("$", "").replace("£", "")
        s = s.replace("€", "").strip()
        try:
            x = float(s)
        except ValueError:
            return None
        if pct:
            x /= 100.0
        return -x if neg else x
    return None


def as_text(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, (dt.datetime, dt.date)):
        return v.isoformat()[:19]
    s = str(v).strip()
    return s or None


def fmt_num(v: float) -> str:
    """Readable for a schema summary: 54880.0 -> '54,880', 0.1234 -> '0.1234'."""
    if v == int(v) and abs(v) < 1e15:
        return f"{int(v):,}"
    return f"{v:,.4f}".rstrip("0").rstrip(".")


def looks_like_label(v) -> bool:
    return isinstance(v, str) and bool(v.strip()) and len(v) <= 60 and as_number(v) is None


# --------------------------------------------------------------------------- #
# Grid
# --------------------------------------------------------------------------- #
class Grid:
    """A sheet as two dense 1-indexed grids: raw values and merge-filled values."""

    def __init__(self, sheet: str, raw: list[list], filled: list[list], forms: list[list]):
        self.sheet = sheet
        self.raw = raw            # [r][c], 0-indexed internally
        self.filled = filled
        self.forms = forms
        self.nrows = len(raw)
        self.ncols = len(raw[0]) if raw else 0

    def v(self, r: int, c: int):
        """1-indexed merge-filled value."""
        if 1 <= r <= self.nrows and 1 <= c <= self.ncols:
            return self.filled[r - 1][c - 1]
        return None


def grid_from_worksheet(ws, ws_f) -> Grid:
    nrows, ncols = ws.max_row or 0, ws.max_column or 0
    raw = [[None] * ncols for _ in range(nrows)]
    forms = [[None] * ncols for _ in range(nrows)]
    for row in ws.iter_rows(min_row=1, max_row=nrows, min_col=1, max_col=ncols):
        for cell in row:
            if cell.value is not None:
                raw[cell.row - 1][cell.column - 1] = cell.value
    if ws_f is not None:
        for row in ws_f.iter_rows(min_row=1, max_row=nrows, min_col=1, max_col=ncols):
            for cell in row:
                v = cell.value
                if isinstance(v, str) and v.startswith("="):
                    forms[cell.row - 1][cell.column - 1] = v
    filled = [r[:] for r in raw]
    for rng in getattr(ws, "merged_cells", []).ranges if hasattr(ws, "merged_cells") else []:
        top = raw[rng.min_row - 1][rng.min_col - 1]
        if top is None:
            continue
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                filled[r - 1][c - 1] = top
    return Grid(ws.title, raw, filled, forms)


def grid_from_csv(path: Path) -> Grid:
    with path.open("r", newline="", encoding="utf-8-sig", errors="replace") as fh:
        sample = fh.read(8192)
        fh.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
            if path.suffix.lower() == ".tsv":
                dialect = csv.excel_tab
        rows = [r for r in csv.reader(fh, dialect)]
    ncols = max((len(r) for r in rows), default=0)
    raw = [[(v if (v is not None and v != "") else None) for v in r] + [None] * (ncols - len(r))
           for r in rows]
    # numbers that were written as text become real numbers downstream via as_number
    return Grid("", raw, [r[:] for r in raw], [[None] * ncols for _ in raw])


# --------------------------------------------------------------------------- #
# Block detection
# --------------------------------------------------------------------------- #
def _runs(sorted_idx: list[int], max_gap: int) -> list[list[int]]:
    """Group indices, splitting where the gap is larger than max_gap."""
    out: list[list[int]] = []
    for i in sorted_idx:
        if out and i - out[-1][-1] <= max_gap + 1:
            out[-1].append(i)
        else:
            out.append([i])
    return out


def find_blocks(g: Grid) -> list[tuple[list[int], list[int]]]:
    """-> [(rows, cols)] 1-indexed. Split at 2+ blank rows, then 2+ blank columns."""
    live_rows = [r for r in range(1, g.nrows + 1)
                 if any(not is_blank(g.v(r, c)) for c in range(1, g.ncols + 1))]
    blocks: list[tuple[list[int], list[int]]] = []
    for row_run in _runs(live_rows, max_gap=1):
        lo, hi = row_run[0], row_run[-1]
        live_cols = [c for c in range(1, g.ncols + 1)
                     if any(not is_blank(g.v(r, c)) for r in range(lo, hi + 1))]
        for col_run in _runs(live_cols, max_gap=1):
            c0, c1 = col_run[0], col_run[-1]
            rows = [r for r in range(lo, hi + 1)
                    if any(not is_blank(g.v(r, c)) for c in range(c0, c1 + 1))]
            if rows:
                blocks.append((rows, list(range(c0, c1 + 1))))
    return blocks


def split_at_headers(g: Grid, rows: list[int], cols: list[int]) -> list[list[int]]:
    """A sheet that stacks several tables one blank row apart is one block to `find_blocks`,
    so every cell below the first table used to carry the first table's headers (a TAM row
    under "2024 estimate" was stored under "Revenue one filing ago", 2026-09-30). A header-ish
    row that follows numbers starts a new table; the title lines just above it (single-cell
    text rows) go with it as its preamble."""
    out: list[list[int]] = []
    cur: list[int] = []
    seen_numbers = False
    for r in rows:
        if cur and seen_numbers and _row_is_headerish(g, r, cols) and not _row_has_numbers(g, r, cols):
            # pull the title rows above into the new table
            k = len(cur)
            while k > 0:
                rr = cur[k - 1]
                vals = [g.v(rr, c) for c in cols]
                nonempty = [v for v in vals if not is_blank(v)]
                if len(nonempty) <= 1 and not _row_has_numbers(g, rr, cols):
                    k -= 1
                else:
                    break
            head, cur = cur[:k], cur[k:]
            if head:
                out.append(head)
            seen_numbers = False
        cur.append(r)
        if _row_has_numbers(g, r, cols) and sum(1 for c in cols if not is_blank(g.v(r, c))) >= 2:
            seen_numbers = True
    if cur:
        out.append(cur)
    return out


def _row_is_headerish(g: Grid, r: int, cols: list[int]) -> bool:
    vals = [g.v(r, c) for c in cols]
    nonempty = [v for v in vals if not is_blank(v)]
    if len(nonempty) < 2:
        return False
    labels = [v for v in nonempty if looks_like_label(v)]
    if len(labels) / len(nonempty) < 0.6:
        return False
    # a value merged across the whole block is a title, not a header band
    if len(set(str(v) for v in nonempty)) == 1 and len(nonempty) >= len(cols) - 1:
        return False
    return True


def _row_has_numbers(g: Grid, r: int, cols: list[int]) -> bool:
    return any(as_number(g.v(r, c)) is not None for c in cols[1:] or cols)


def detect_header(g: Grid, rows: list[int], cols: list[int]) -> tuple[list[int], int]:
    """-> (header_rows, first_data_row). Walks back from the first numeric row."""
    first_num = None
    for r in rows:
        vals = [g.v(r, c) for c in cols]
        if sum(1 for v in vals if not is_blank(v)) >= 2 and _row_has_numbers(g, r, cols):
            first_num = r
            break
    if first_num is None:                      # all-text block: first row is the header
        return [rows[0]], (rows[1] if len(rows) > 1 else rows[0] + 1)
    header: list[int] = []
    r = first_num - 1
    while r >= rows[0] and len(header) < MAX_HEADER_ROWS and _row_is_headerish(g, r, cols):
        header.insert(0, r)
        r -= 1
    if not header:
        header = [first_num - 1] if first_num - 1 >= rows[0] else [first_num]
    return header, first_num


def combine_labels(parts: list[str]) -> str:
    """'FY25 actual' over 'Q3' -> 'FY25 actual Q3'; drops repeats and contained parts."""
    parts = [p.strip() for p in parts if p and p.strip()]
    if not parts:
        return ""
    keep: list[str] = []
    for p in parts:
        if any(p.lower() in k.lower() for k in keep):
            continue
        keep = [k for k in keep if k.lower() not in p.lower()] + [p]
    words, seen, out = " ".join(keep).split(), set(), []
    for w in words:
        if w.lower() in seen:
            continue
        seen.add(w.lower())
        out.append(w)
    return " ".join(out)


# --------------------------------------------------------------------------- #
# Block -> table
# --------------------------------------------------------------------------- #
class TableSpec:
    def __init__(self, **kw):
        self.__dict__.update(kw)


def build_table(g: Grid, rows: list[int], cols: list[int], file: str,
                taken: set[str], name_hint: str) -> TableSpec | None:
    header_rows, first_data = detect_header(g, rows, cols)
    data_rows = [r for r in rows if r >= first_data]
    if len(data_rows) < 1 or len(cols) < 2:
        return None

    # preamble read from the raw grid, so a title merged across the sheet appears once
    pre: list[str] = []
    for r in rows:
        if r >= header_rows[0]:
            break
        for c in cols:
            v = g.raw[r - 1][c - 1]
            if not is_blank(v):
                pre.append(str(v).strip())
    title = " / ".join(pre)[:300]

    labels = [combine_labels([as_text(g.v(hr, c)) or "" for hr in header_rows])
              or f"col {col_letter(c)}" for c in cols]
    names, seen = [], set()
    for lab, c in zip(labels, cols):
        names.append(unique_ident(ident(lab, f"col_{col_letter(c)}"), seen))

    grid_vals = [[g.v(r, c) for c in cols] for r in data_rows]

    types = []
    for j in range(len(cols)):
        vals = [row[j] for row in grid_vals if not is_blank(row[j])]
        nums = sum(1 for v in vals if as_number(v) is not None)
        if vals and nums / len(vals) >= 0.8:
            types.append("REAL")
        elif any(isinstance(v, (dt.datetime, dt.date)) for v in vals):
            types.append("DATE")
        else:
            types.append("TEXT")
    if types and types[0] == "REAL" and all(looks_like_label(row[0]) for row in grid_vals[:3] if not is_blank(row[0])):
        types[0] = "TEXT"

    out_rows = []
    for r, gv in zip(data_rows, grid_vals):
        rec = []
        for j, v in enumerate(gv):
            if types[j] == "REAL":
                rec.append(as_number(v))
            else:
                rec.append(as_text(v))
        label0 = as_text(gv[0]) or ""
        rec.append(r)
        rec.append(1 if TOTAL_RE.search(label0) else 0)
        out_rows.append(rec)

    total_rows = [str(r) for r, rec in zip(data_rows, out_rows) if rec[-1]]
    base = ident(name_hint)
    name = unique_ident(base, taken)
    return TableSpec(
        name=name, file=file, sheet=g.sheet, title=title,
        header_rows=header_rows, first_data_row=data_rows[0], last_data_row=data_rows[-1],
        cols=cols, labels=labels, names=names, types=types, rows=out_rows,
        total_rows=",".join(total_rows), first_col_letter=col_letter(cols[0]),
    )


def create_and_fill(con, t: TableSpec) -> None:
    decls = ", ".join(f'"{n}" {("REAL" if ty == "REAL" else "TEXT")}' for n, ty in zip(t.names, t.types))
    con.execute(f'DROP TABLE IF EXISTS "{t.name}"')
    con.execute(f'CREATE TABLE "{t.name}" ({decls}, "_row" INTEGER, "_is_total" INTEGER)')
    ph = ",".join("?" * (len(t.names) + 2))
    con.executemany(f'INSERT INTO "{t.name}" VALUES ({ph})', t.rows)
    con.execute("DELETE FROM _tables WHERE name=?", (t.name,))
    con.execute(
        "INSERT INTO _tables VALUES (?,?,?,?,?,?,?,?,?,?,?)",
        (t.name, t.file, t.sheet, t.title, ",".join(map(str, t.header_rows)),
         t.first_data_row, t.last_data_row, t.first_col_letter,
         len(t.rows), len(t.names), t.total_rows),
    )
    con.execute("DELETE FROM _columns WHERE table_name=?", (t.name,))
    for j, (n, lab, c, ty) in enumerate(zip(t.names, t.labels, t.cols, t.types)):
        vals = [r[j] for r in t.rows]
        nn = sum(1 for v in vals if v is None)
        samples = [v for v in vals if v is not None][:3]
        samp = " | ".join(
            (fmt_num(v) if isinstance(v, float) else str(v))[:40] for v in samples
        )
        con.execute("INSERT INTO _columns VALUES (?,?,?,?,?,?,?,?)",
                    (t.name, j, n, lab, col_letter(c), ty, nn, samp))


# --------------------------------------------------------------------------- #
# The universal long table
# --------------------------------------------------------------------------- #
def emit_cells(con, g: Grid, file: str, tables: list[TableSpec], is_csv: bool) -> tuple[int, int]:
    """Every non-empty cell -> cells(...). Returns (n_cells, n_null_formulas)."""
    # header label / row label per (row, col), from whichever table owns the cell
    header_at: dict[int, str] = {}                 # column fallback, rows no table owns
    header_rc: dict[tuple[int, int], str] = {}     # the table that owns the row wins
    row_label_at: dict[int, str] = {}
    for t in tables:
        for c, lab in zip(t.cols, t.labels):
            header_at.setdefault(c, lab)
            for r in range(t.header_rows[0], t.last_data_row + 1):
                header_rc.setdefault((r, c), lab)
        lab_col = t.cols[0]
        for r in range(t.first_data_row, t.last_data_row + 1):
            lab = as_text(g.v(r, lab_col))
            if lab:
                row_label_at[r] = lab

    n_rows_total = sum(len(t.rows) for t in tables) or g.nrows
    if is_csv and n_rows_total > CELLS_MAX_DATA_ROWS:
        return 0, 0

    batch, n_null = [], 0
    for r in range(1, g.nrows + 1):
        for c in range(1, g.ncols + 1):
            raw = g.raw[r - 1][c - 1]
            formula = g.forms[r - 1][c - 1]
            if is_blank(raw) and not formula:
                continue
            num = as_number(raw)
            txt = None if num is not None else as_text(raw)
            if formula and num is None and txt is None:
                n_null += 1
            batch.append((
                file, g.sheet, r, c, f"{col_letter(c)}{r}",
                header_rc.get((r, c), header_at.get(c)), row_label_at.get(r),
                num, txt, 1 if formula else 0, formula,
            ))
    con.executemany("INSERT INTO cells VALUES (?,?,?,?,?,?,?,?,?,?,?)", batch)
    return len(batch), n_null


# --------------------------------------------------------------------------- #
# Entry point
# --------------------------------------------------------------------------- #
def ingest_spreadsheet(con, path: Path, rel: str, taken: set[str]) -> dict:
    """Ingest one workbook or csv. Returns {'tables': [...], 'n_cells': int, ...}."""
    import openpyxl

    stem = Path(rel).stem
    grids: list[Grid] = []
    if path.suffix.lower() in CSV_EXT:
        grids.append(grid_from_csv(path))
    else:
        wb = openpyxl.load_workbook(path, data_only=True)
        try:
            wb_f = openpyxl.load_workbook(path, data_only=False)
        except Exception:
            wb_f = None
        for ws in wb.worksheets:
            if ws.sheet_state != "visible":
                continue
            ws_f = wb_f[ws.title] if wb_f is not None and ws.title in wb_f.sheetnames else None
            grids.append(grid_from_worksheet(ws, ws_f))

    made: list[TableSpec] = []
    n_cells = n_null = 0
    for g in grids:
        if g.nrows == 0 or g.ncols == 0:
            continue
        blocks = [(sub, cols) for rows, cols in find_blocks(g) for sub in split_at_headers(g, rows, cols)]
        sheet_tables: list[TableSpec] = []
        base_hint = stem if not g.sheet else g.sheet
        for i, (rows, cols) in enumerate(blocks):
            hint = base_hint
            if i > 0:
                # second table on the same sheet: name it after its own first column
                hrows, _ = detect_header(g, rows, cols)
                lab = as_text(g.v(hrows[-1], cols[0])) or f"block{i + 1}"
                hint = f"{base_hint}_{lab}"
            t = build_table(g, rows, cols, rel, taken, hint)
            if t is not None:
                create_and_fill(con, t)
                sheet_tables.append(t)
                made.append(t)
        c, nf = emit_cells(con, g, rel, sheet_tables, is_csv=not g.sheet)
        n_cells += c
        n_null += nf
    return {"tables": made, "n_cells": n_cells, "n_null_formulas": n_null}


def drop_file_tables(con, rel: str) -> None:
    for (name,) in con.execute("SELECT name FROM _tables WHERE file=?", (rel,)).fetchall():
        con.execute(f'DROP TABLE IF EXISTS "{name}"')
        con.execute("DELETE FROM _columns WHERE table_name=?", (name,))
    con.execute("DELETE FROM _tables WHERE file=?", (rel,))
    con.execute("DELETE FROM cells WHERE file=?", (rel,))
