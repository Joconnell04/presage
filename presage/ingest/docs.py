"""Documents -> citable chunks.

pdf via pymupdf (page), docx via python-docx (heading + tables), pptx via
python-pptx (slide number, title, body, tables, speaker notes), txt/md by heading.

Every chunk carries a `loc` that is ready to print in a citation:
"page 4", "slide 7 'Risks and asks'", "section Churn and retention".
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

PDF_EXT = {".pdf"}
DOCX_EXT = {".docx"}
PPTX_EXT = {".pptx"}
TEXT_EXT = {".txt", ".md", ".markdown", ".text"}
DOC_EXT = PDF_EXT | DOCX_EXT | PPTX_EXT | TEXT_EXT

TARGET_WORDS = 300          # aim
MIN_WORDS = 160             # do not flush below this unless the unit ends
MAX_WORDS = 420             # hard stop
OVERLAP_SENTENCES = 2


@dataclass
class Chunk:
    id: int
    file: str
    loc: str
    text: str
    score: float = 0.0
    kind: str = ""
    unit: int = 0
    heading: str = ""
    context: str = ""     # deep pass: 50-100 token blurb, indexed but never displayed

    def cite(self) -> str:
        return f"{self.file} › {self.loc}"

    def indexed(self) -> str:
        """What BM25 and the dense model see. Display uses `text`."""
        return f"{self.context}\n{self.text}" if self.context else self.text


@dataclass
class DocResult:
    chunks: list[Chunk] = field(default_factory=list)
    headings: list[str] = field(default_factory=list)
    n_units: int = 0
    unit: str = ""
    kind: str = ""


_SENT = re.compile(r"(?<=[.!?:])\s+(?=[A-Z0-9•\-])|\n{2,}")


def split_sentences(text: str) -> list[str]:
    parts = [p.strip() for p in _SENT.split(text) if p and p.strip()]
    return parts or ([text.strip()] if text.strip() else [])


def chunk_text(text: str, file: str, loc: str, kind: str, unit: int,
               heading: str, start_id: int) -> list[Chunk]:
    """~300 words per chunk, broken on sentences, 1-2 sentences of overlap."""
    sents = split_sentences(text)
    out: list[Chunk] = []
    buf: list[str] = []
    n = 0
    for s in sents:
        w = len(s.split())
        if buf and n + w > MAX_WORDS and n >= MIN_WORDS:
            out.append(Chunk(start_id + len(out), file, loc, " ".join(buf), kind=kind,
                             unit=unit, heading=heading))
            buf = buf[-OVERLAP_SENTENCES:]
            n = sum(len(x.split()) for x in buf)
        buf.append(s)
        n += w
        if n >= TARGET_WORDS:
            out.append(Chunk(start_id + len(out), file, loc, " ".join(buf), kind=kind,
                             unit=unit, heading=heading))
            buf = buf[-OVERLAP_SENTENCES:]
            n = sum(len(x.split()) for x in buf)
    tail = " ".join(buf).strip()
    if tail and (not out or len(tail.split()) > sum(len(x.split()) for x in buf[:OVERLAP_SENTENCES])):
        out.append(Chunk(start_id + len(out), file, loc, tail, kind=kind, unit=unit, heading=heading))
    return out


# --------------------------------------------------------------------------- #
def read_pdf(path: Path, rel: str, start_id: int) -> DocResult:
    import pymupdf

    res = DocResult(kind="pdf", unit="page")
    doc = pymupdf.open(path)
    res.n_units = doc.page_count
    nid = start_id
    for i, page in enumerate(doc, start=1):
        text = page.get_text("text").strip()
        if not text:
            continue
        first = next((ln.strip() for ln in text.splitlines() if ln.strip()), "")
        if 3 <= len(first) <= 90:
            res.headings.append(first)
        cs = chunk_text(text, rel, f"page {i}", "pdf", i, first, nid)
        nid += len(cs)
        res.chunks.extend(cs)
    doc.close()
    return res


def read_docx(path: Path, rel: str, start_id: int) -> DocResult:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    res = DocResult(kind="docx", unit="section")
    doc = Document(str(path))
    body = doc.element.body
    heading, parts = "", []
    sections: list[tuple[str, str]] = []

    def flush():
        if parts:
            sections.append((heading, "\n".join(parts)))

    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            p = Paragraph(child, doc)
            txt = p.text.strip()
            if not txt:
                continue
            style = (p.style.name or "") if p.style is not None else ""
            if style.startswith("Heading") or style == "Title":
                flush()
                heading, parts = txt, []
                res.headings.append(txt)
            else:
                parts.append(txt)
        elif tag == "tbl":
            t = Table(child, doc)
            rows = [[c.text.strip() for c in r.cells] for r in t.rows]
            if not rows:
                continue
            head = rows[0]
            lines = ["Table: " + " | ".join(head)]
            for r in rows[1:]:
                lines.append("; ".join(f"{h}: {v}" for h, v in zip(head, r) if v))
            parts.append("\n".join(lines))
    flush()
    res.n_units = len(sections)
    nid = start_id
    for i, (h, text) in enumerate(sections, start=1):
        loc = f"section {h}" if h else f"section {i}"
        cs = chunk_text((h + "\n" + text) if h else text, rel, loc, "docx", i, h, nid)
        nid += len(cs)
        res.chunks.extend(cs)
    return res


def read_pptx(path: Path, rel: str, start_id: int) -> DocResult:
    from pptx import Presentation

    res = DocResult(kind="pptx", unit="slide")
    prs = Presentation(str(path))
    nid = start_id
    slides = list(prs.slides)
    res.n_units = len(slides)
    for i, slide in enumerate(slides, start=1):
        title = ""
        try:
            if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
                title = slide.shapes.title.text.strip()
        except (AttributeError, ValueError):
            title = ""
        body: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                txt = shape.text_frame.text.strip()
                if txt and txt != title:
                    body.append(txt)
            if getattr(shape, "has_table", False):
                tbl = shape.table
                rows = [[c.text.strip() for c in r.cells] for r in tbl.rows]
                if rows:
                    head = rows[0]
                    body.append("Table: " + " | ".join(head))
                    for r in rows[1:]:
                        body.append("; ".join(f"{h}: {v}" for h, v in zip(head, r) if v))
        if not title and body:
            first = body[0].splitlines()[0]
            title = first[:80]
            if body[0].strip() == title:          # a title in a plain textbox, not a placeholder
                body.pop(0)
        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        if title:
            res.headings.append(title)
        loc = f"slide {i} '{title}'" if title else f"slide {i}"
        text = "\n".join([title] + body).strip()
        if text:
            cs = chunk_text(text, rel, loc, "pptx", i, title, nid)
            nid += len(cs)
            res.chunks.extend(cs)
        if notes:
            cs = chunk_text("Speaker notes: " + notes, rel, loc + " notes", "pptx-notes", i, title, nid)
            nid += len(cs)
            res.chunks.extend(cs)
    return res


_MD_H = re.compile(r"^(#{1,6})\s+(.*\S)\s*$")


def read_text(path: Path, rel: str, start_id: int) -> DocResult:
    res = DocResult(kind=path.suffix.lstrip(".").lower() or "txt", unit="section")
    raw = path.read_text(encoding="utf-8", errors="replace")
    sections: list[tuple[str, list[str]]] = [("", [])]
    for line in raw.splitlines():
        m = _MD_H.match(line)
        if m:
            sections.append((m.group(2), []))
            res.headings.append(m.group(2))
        else:
            sections[-1][1].append(line)
    sections = [(h, "\n".join(b).strip()) for h, b in sections if h or "\n".join(b).strip()]
    res.n_units = len(sections)
    nid = start_id
    for i, (h, text) in enumerate(sections, start=1):
        loc = f"section {h}" if h else f"section {i}"
        cs = chunk_text((h + "\n" + text) if h else text, rel, loc, "text", i, h, nid)
        nid += len(cs)
        res.chunks.extend(cs)
    return res


def read_document(path: Path, rel: str, start_id: int) -> DocResult:
    ext = path.suffix.lower()
    if ext in PDF_EXT:
        return read_pdf(path, rel, start_id)
    if ext in DOCX_EXT:
        return read_docx(path, rel, start_id)
    if ext in PPTX_EXT:
        return read_pptx(path, rel, start_id)
    return read_text(path, rel, start_id)
